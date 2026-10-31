from docx import Document
from pathlib import Path
from docx.shared import Pt
R=Path('/Users/rafaelyudianto/Documents/COMP8035041-Assignment-II')
O=R/'Assignment-II-DataCo-Laporan-Final-High-Quality.docx'
d=Document(O)
def find(t):
    for p in d.paragraphs:
        if p.text.strip().startswith(t):return p
    raise KeyError(t[:40])
p=find('Urutan sentralitas pada proyeksi kategori konsisten')
new=d.add_paragraph()
new.add_run('Sepuluh kategori dengan degree tertinggi dipilih dari 51 kategori yang muncul pada graf; pemilihan ini hanya untuk visualisasi dan tidak mengubah statistik graf lengkap. Ketiadaan sinyal pada uji kategori saja (AP 0,576630; ROC-AUC 0,500000) konsisten dengan temuan bahwa struktur kategori tidak memisahkan keterlambatan, meskipun tingkat keterlambatan per kategori bervariasi antara 0,4963 dan 0,6885. Variasi tersebut tidak cukup untuk menghasilkan pemeringkatan pesanan yang baik, sehingga metrik graf belum layak dijadikan fitur tanpa pengujian tambahan.')
for r in new.runs:r.font.size=Pt(11)
p._p.addnext(new._p)
d.save(O)
print('insight 6.2 ditambahkan')
