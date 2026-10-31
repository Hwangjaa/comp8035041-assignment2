from pathlib import Path
from docx import Document
import fitz, re, json, hashlib
R=Path('/Users/rafaelyudianto/Documents/COMP8035041-Assignment-II')
O=R/'Assignment-II-DataCo-Laporan-Final-High-Quality.docx'
PDF=Path('/Users/rafaelyudianto/.hermes/cache/scratch/dataco-hq/Assignment-II-DataCo-Laporan-Final-High-Quality.pdf')
b=Document(O);pdf=fitz.open(PDF)
norm=lambda x:re.sub(r'\s+','',x).casefold()
pt=[norm(x.get_text()) for x in pdf]
entries=[];bad=[]
for p in b.paragraphs:
    if '\t' in p.text and p.text.rsplit('\t',1)[-1].isdigit():
        t,n=p.text.rsplit('\t',1);n=int(n)
        if 0<n<=len(pt) and norm(t) in pt[n-1]:entries.append((t,n))
        else:bad.append((t,n))
empty=[i+1 for i,pg in enumerate(pdf) if len(pg.get_text().strip())<10]
geo=[]
for i,pg in enumerate(pdf):
    for bl in pg.get_text('dict')['blocks']:
        x0,y0,x1,y1=bl['bbox']
        if min(x0,y0)<-0.5 or x1>pg.rect.width+0.5 or y1>pg.rect.height+0.5:geo.append((i+1,tuple(round(v,1) for v in bl['bbox'])))
codes=[c for c in re.findall(r'\b(?:D2|H1|M0|F[2-8]|R0|T0|U0|G1)\b','\n'.join(p.text for p in b.paragraphs))]
rep={'docx':str(O),'pdf':str(PDF),'pages':len(pdf),'sha256':hashlib.sha256(O.read_bytes()).hexdigest(),
 'toc_entries':len(entries),'toc_bad':bad,'empty_pages':empty,'geometry_issues':geo[:5],
 'internal_codes':codes[:10],'tables':len(b.tables),'images':len(b.inline_shapes),
 'headings':len([p for p in b.paragraphs if p.style.name.startswith('Heading')]),
 'under_40':len(pdf)<=40}
json.dump(rep,open(R/'analisis-lanjutan/FINAL-VERIFY.json','w'),indent=2)
print(json.dumps(rep,indent=2,ensure_ascii=False))
