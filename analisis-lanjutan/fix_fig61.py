from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches
R=Path('/Users/rafaelyudianto/Documents/COMP8035041-Assignment-II')
OUT=R/'analisis-lanjutan';SRC=R/'Assignment-II-DataCo-Laporan-Final-High-Quality.docx'
d=Document(SRC)
# buat rId baru via run sementara
tmp=d.add_paragraph();run=tmp.add_run();run.add_picture(str(OUT/'D_graph_subset_fixed.png'),width=Inches(6))
new_rId=run._r.find('.//'+qn('a:blip')).get(qn('r:embed'))
tmp._p.getparent().remove(tmp._p)
print('new rId:',new_rId)
body=list(d.element.body);done=False
for i,el in enumerate(body):
    if el.tag==qn('w:p') and el.findall('.//'+qn('w:drawing')):
        nxt=body[i+1] if i+1<len(body) else None
        cap=''.join(t.text or '' for t in nxt.iter(qn('w:t'))) if nxt is not None and nxt.tag==qn('w:p') else ''
        if cap.startswith('Gambar 6.1 Subset graf'):
            for b in el.findall('.//'+qn('a:blip')):b.set(qn('r:embed'),new_rId)
            done=True;print('replaced Gambar 6.1 drawing at body idx',i)
d.save(SRC)
d2=Document(SRC);body2=list(d2.element.body)
for i,el in enumerate(body2):
    if el.tag==qn('w:p') and el.findall('.//'+qn('w:drawing')):
        nxt=body2[i+1] if i+1<len(body2) else None
        cap=''.join(t.text or '' for t in nxt.iter(qn('w:t'))) if nxt is not None and nxt.tag==qn('w:p') else ''
        print('now:',cap[:40],'->',el.find('.//'+qn('a:blip')).get(qn('r:embed')))
print('done:',done)
