from pathlib import Path
from copy import deepcopy
from docx import Document
R=Path('/Users/rafaelyudianto/Documents/COMP8035041-Assignment-II')
O=R/'Assignment-II-DataCo-Laporan-Final-High-Quality.docx'
d=Document(O)
new='Pelatihan dan evaluasi akhir menghasilkan model, probabilitas, prediksi, serta statistik graf yang tersimpan sebagai artefak eksekusi. Spesifikasi fitur, parameter model, lingkungan percobaan, dan ambang klasifikasi tercatat lengkap sehingga setiap angka hasil dapat ditelusuri. Logistic Regression scikit-learn mencatat 415 iterasi dan Spark MLlib 71 iterasi. Jumlah iterasi bukan ukuran kecepatan karena konfigurasi optimisasi serta regularisasinya berbeda; benchmark waktu dan skalabilitas multi-node tidak dilakukan.'
done=0
for p in d.paragraphs:
    if 'Kontrak fitur, parameter model' in p.text:
        first=next((r for r in p.runs if r.text),None)
        props=deepcopy(first._r.rPr) if first is not None and first._r.rPr is not None else None
        p.clear();r=p.add_run(new)
        if props is not None:r._r.insert(0,props)
        r.italic=False;done+=1
d.save(O)
print('paragraphs fixed:',done)
d2=Document(O)
t='\n'.join(p.text for p in d2.paragraphs)
print('kontrak fitur remains:','kontrak fitur' in t.lower())
print('feature-contract.json remains:','feature-contract.json' in t)
