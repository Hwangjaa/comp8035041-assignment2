"""Perbaiki label TOC entry (izinkan elemen w:tab di prefix run)."""
from docx import Document
from docx.oxml.ns import qn
from pathlib import Path
R=Path('/Users/rafaelyudianto/Documents/COMP8035041-Assignment-II')
O=R/'Assignment-II-DataCo-Laporan-Final-High-Quality.docx'
d=Document(O)
ALLOWED={qn('w:rPr'),qn('w:t'),qn('w:tab')}
def rename_toc(old,new):
    cnt=0
    for p in d.paragraphs:
        if '\t' not in p.text or old not in p.text:continue
        prefix=[]
        for r in p.runs:
            if any(x.tag not in ALLOWED for x in r._r):break
            prefix.append(r)
        joined=''.join(r.text for r in prefix)
        if old in joined:
            prefix[0].text=joined.replace(old,new)
            for r in prefix[1:]:r.text=''
            cnt+=1
    return cnt
a=rename_toc('Tabel 5.5 Hasil evaluasi pada data validation','Tabel 5.6 Hasil evaluasi pada data validation')
b=rename_toc('Tabel 5.4 Hasil evaluasi pada data test','Tabel 5.5 Hasil evaluasi pada data test')
d.save(O)
print('toc validation->5.6:',a,'toc test->5.5:',b)
d2=Document(O);on=False
seq=[]
for p in d2.paragraphs:
    t=p.text.strip()
    if t=='DAFTAR TABEL':on=True;continue
    if t=='DAFTAR GAMBAR':break
    if on and '\t' in p.text:
        lbl,n=p.text.rsplit('\t',1);seq.append(int(n));print('  ',lbl,'>>',n)
print('sorted:',seq==sorted(seq))
