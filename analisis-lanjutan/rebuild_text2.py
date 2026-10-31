"""Tulis ulang 5.6 & 6.2, perbarui Tabel 6.1/6.4, tambah tabel komunitas & komponen, ganti gambar 5.1/6.2, tabel jadwal."""
import json
from pathlib import Path
from copy import deepcopy
import numpy as np, pandas as pd
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
R=Path('/Users/rafaelyudianto/Documents/COMP8035041-Assignment-II')
OUT=R/'analisis-lanjutan';SRC=R/'Assignment-II-DataCo-Laporan-Final-High-Quality.docx'
F=json.load(open(OUT/'F_main.json'));G=json.load(open(OUT/'G_ablation_schedule.json'));H=json.load(open(OUT/'H_graph.json'))
def idn(x,n=6):return f'{x:.{n}f}'.replace('.',',')
def n2(x):return f'{x:,}'.replace(',','.')
d=Document(SRC)
# late rate range semua kategori
edges=pd.read_csv(R/'eksekusi-final-lokal/run-actual/graph-edges.csv')
orders=pd.read_csv(R/'eksekusi-final-lokal/run-actual/orders.csv')
lab=orders.set_index('Order Id')['Late_delivery_risk'];e2=edges.copy();e2['late']=e2['order_id'].map(lab)
g=e2.groupby('category_id')['late'].agg(['mean','count'])
cat_min=float(g['mean'].min());cat_max=float(g['mean'].max());cat_min_id=int(g['mean'].idxmin());cat_max_id=int(g['mean'].idxmax())
overall=float(lab.mean())
def find_para(t):
    for p in d.paragraphs:
        if p.text.strip().startswith(t) and '\t' not in p.text:return p
    raise KeyError(t[:40])
def set_text(p,new):
    first=next((r for r in p.runs if r.text),None);props=deepcopy(first._r.rPr) if (first is not None and first._r.rPr is not None) else None
    p.clear();r=p.add_run(new)
    if props is not None:r._r.insert(0,props)
    r.italic=False
def add_para_after(anchor,text,style=None):
    p=d.add_paragraph();p.add_run(text)
    if style:p.style=d.styles[style]
    anchor._p.addnext(p._p);return p
def add_table_after(anchor,rows,style_tbl):
    t=d.add_table(rows=len(rows),cols=len(rows[0]))
    try:t.style=style_tbl
    except Exception:pass
    for i,row in enumerate(rows):
        for j,v in enumerate(row):
            c=t.cell(i,j);c.text=str(v)
            for pp in c.paragraphs:
                for rr in pp.runs:rr.font.size=Pt(8.5);rr.bold=(i==0)
    anchor._p.addnext(t._tbl);return t
def tbl_by_caption(cap):
    from docx.table import Table
    p=find_para(cap);nxt=p._p.getnext()
    while nxt is not None and nxt.tag!=qn('w:tbl'):nxt=nxt.getnext()
    return Table(nxt,d)
def rebuild(tbl,rows):
    while len(tbl.rows)>len(rows):tbl._tbl.remove(tbl.rows[-1]._tr)
    while len(tbl.rows)<len(rows):tbl.add_row()
    for i,row in enumerate(rows):
        for j,v in enumerate(row):
            if j<len(tbl.columns):
                c=tbl.cell(i,j);c.text=str(v)
                for pp in c.paragraphs:
                    for rr in pp.runs:rr.font.size=Pt(8.5);rr.bold=(i==0)
tb=d.tables[16].style
# ====== 5.6 ditulis ulang ======
t=F['main']['test'];bs=F['bootstrap']['sklearn_lr']
set_text(find_para('Pada data test, AP baseline adalah'),
 f"Pada data test, Average Precision baseline adalah {idn(t['baseline']['ap'])} sedangkan kedua implementasi Logistic Regression menghasilkan {idn(t['sklearn_lr']['ap'])}. ROC-AUC baseline {idn(t['baseline']['auc'])} dan kedua implementasi {idn(t['sklearn_lr']['auc'])}. Karena regularisasi disetarakan dan seed sama, scikit-learn dan Spark MLlib menghasilkan probabilitas yang identik secara numerik, sehingga selisih AP dan ROC-AUC keduanya nol. Pada validation, AP {idn(F['main']['validation']['sklearn_lr']['ap'])} dan ROC-AUC {idn(F['main']['validation']['sklearn_lr']['auc'])}. Interval kepercayaan bootstrap 95% untuk AP adalah [{idn(bs['ap']['lo95'])}, {idn(bs['ap']['hi95'])}] dan untuk ROC-AUC [{idn(bs['auc']['lo95'])}, {idn(bs['auc']['hi95'])}].")
set_text(find_para('Pada ambang yang dipilih dari validation'),
 "Pada ambang F1 dari validation, seluruh 8.043 pesanan test diprediksi sebagai kelas 1 oleh baseline maupun kedua implementasi. Confusion matrix menunjukkan TN=0, FP=3.407, FN=0, dan TP=4.636. Akibatnya precision sama dengan proporsi positif test (0,576402), recall 1,000000, dan binary F1 0,731288. Kriteria F1 pada validation karena itu tidak menghasilkan aturan yang mengalahkan kebijakan semua positif: nilai F1 identik dengan baseline. Optimasi F1 pada validation tidak menjamin keputusan yang informatif pada test.")
set_text(find_para('Karena seluruh pesanan test dinyatakan positif'),
 "Dengan aturan peringkat, 20% pesanan dengan skor tertinggi (1.609 pesanan) memuat 1.398 pesanan terlambat sehingga precision@20% 0,868863, recall 0,301553, dan lift 1,507 terhadap prevalensi. Pada 30% (2.413 pesanan), precision 0,854952 dan recall 0,444996. Presisi pada bagian teratas peringkat meningkat dari 0,576 ke sekitar 0,85 sampai 0,87, tetapi recall tetap terbatas pada 30 sampai 45 persen. Nilai ini menggambarkan kemampuan pemeringkatan, bukan keputusan operasional pada satu ambang.")
set_text(find_para('False positive adalah pesanan tidak terlambat'),
 "Uji ablation menunjukkan hampir seluruh sinyal berasal dari Days for shipment (scheduled). Dengan fitur itu saja, AP 0,724911 dan ROC-AUC 0,715880. Tanpa fitur tersebut, AP turun ke 0,583607 dan ROC-AUC 0,507305, mendekati acak. Kategori saja menghasilkan AP 0,576630 dengan ROC-AUC 0,500000; nilai AUC 0,500000 bukan berarti model konstan karena probabilitasnya memiliki 62 nilai berbeda, melainkan urutan skornya tidak memisahkan kelas. Segmen dan market saja menghasilkan AP 0,581831. Model non-linear menaikkan AP menjadi 0,820887 pada Random Forest dan 0,821702 pada Gradient Boosting, dibandingkan 0,736424 pada Logistic Regression dengan fitur yang sama. Selisih ini menunjukkan batas utama terletak pada bentuk hubungan yang linear, bukan pada ketersediaan fitur.")
set_text(find_para('Perbandingan implementasi menggunakan data'),
 "Uji spesifikasi jadwal menjelaskan sumber sinyal tersebut. Tingkat keterlambatan berbeda tajam antarnilai jadwal: 0 hari sebesar 0,4837, 1 hari 1,0000, 2 hari 0,7999, dan 4 hari 0,3985. Hubungan ini tidak monoton sehingga jadwal sebagai variabel numerik tunggal tidak sesuai bentuknya. Ketika jadwal diubah menjadi representasi one-hot, AP naik ke 0,821496 dengan interval kepercayaan 95% [0,813022, 0,830012], hampir menyamai Gradient Boosting. Dengan demikian Logistic Regression dengan jadwal numerik mengalami kesalahan spesifikasi, bukan kekurangan fitur. Varian one-hot disajikan sebagai eksperimen tambahan; penggantian desain utama memerlukan keputusan pemilik penelitian.")
set_text(find_para('Generalisasi dibatasi oleh seleksi pesanan nonbatal'),
 "Dengan demikian nilai operasional belum terbukti. Kriteria F1 tidak mengalahkan aturan semua positif, dan sebagian besar kemampuan model berasal dari satu atribut jadwal yang berkaitan langsung dengan mode pengiriman. Evaluasi ini bersifat retrospektif dan tidak membuktikan pengurangan keterlambatan. Generalisasi juga dibatasi oleh seleksi pesanan nonbatal berdasarkan status akhir, order_date sebagai proksi kronologi, dan waktu tersedianya label yang tidak diketahui. Data tidak memungkinkan rekonstruksi nilai fitur tepat pada waktu prediksi, sehingga validasi operasional memerlukan pencatatan waktu ketersediaan fitur dan label serta evaluasi prospektif yang terpisah.")
# ====== Tabel 6.1 tambah baris ======
rebuild(tbl_by_caption('Tabel 6.1 Definisi graf pesanan dan kategori'),[
 ['Komponen','Definisi'],
 ['Node pesanan','Satu Order Id yang dipertahankan pada kohort nonbatal'],
 ['Node kategori','Satu Category Id yang muncul pada pesanan tersebut'],
 ['Edge bipartit','Pasangan unik pesanan-kategori tanpa bobot dan tanpa arah'],
 ['Proyeksi kategori','Graf kategori-kategori; bobot edge = jumlah pesanan yang memuat kedua kategori'],
 ['Weighted degree','Jumlah bobot edge yang bersinggungan dengan suatu kategori pada proyeksi'],
 ['PageRank','Sentralitas berbobot dengan alpha=0,85 pada proyeksi'],
 ['Eigenvector centrality','Sentralitas vektor eigen pada proyeksi, toleransi 1e-12'],
 ['Betweenness centrality','Proporsi lintasan terpendek yang melewati kategori; dihitung eksak pada 33 node proyeksi'],
 ['Komunitas','Pengelompokan greedy modularity pada proyeksi, resolusi 1,0'],
 ['Tingkat keterlambatan','Proporsi pesanan berlabel 1 yang terhubung dengan kategori atau komponen']])
# ====== Tabel 6.4 betweenness eksak ======
cp=H
D=json.load(open(OUT/'D_graph.json'))
rebuild(tbl_by_caption('Tabel 6.4 Sentralitas proyeksi kategori-kategori'),[
 ['Metrik','Lima teratas'],
 ['Weighted degree',', '.join(f"{x['category']} ({n2(x['weighted_degree'])})" for x in json.load(open(OUT/'D_graph.json'))['category_projection']['weighted_degree_top10'][:5])],
 ['PageRank (alpha=0,85)',', '.join(f"{x['category']} ({idn(x['pagerank'])})" for x in json.load(open(OUT/'D_graph.json'))['category_projection']['pagerank_top10'][:5])],
 ['Eigenvector',', '.join(f"{x['category']} ({idn(x['eigenvector'])})" for x in json.load(open(OUT/'D_graph.json'))['category_projection']['eigenvector_top10'][:5])],
 ['Betweenness (eksak, 33 node)',', '.join(f"{x['category']} ({idn(x['betweenness'])})" for x in cp['betweenness_exact']['top10'][:5])],
 ['Komunitas (greedy modularity)',f"{cp['communities']['n']} komunitas, ukuran {cp['communities']['sizes']}"]])
# ====== Komunitas + komponen kecil + jadwal (berurutan) ======
def cap_after(el,text):
    p=d.add_paragraph(text);p.style=d.styles['Table Caption'];el.addnext(p._p);return p
def table_after(el,rows):
    t=d.add_table(rows=len(rows),cols=len(rows[0]))
    try:t.style=tb
    except Exception:pass
    for i,row in enumerate(rows):
        for j,v in enumerate(row):
            c=t.cell(i,j);c.text=str(v)
            for pp in c.paragraphs:
                for rr in pp.runs:rr.font.size=Pt(8.5);rr.bold=(i==0)
    el.addnext(t._tbl);return t
a=find_para('Komponen terhubung terbesar memuat')
c66=cap_after(a._p,'Tabel 6.6 Anggota dua komunitas pada proyeksi kategori')
rows=[['Komunitas','Jumlah kategori','Anggota (Category Id)']]
for i,m in enumerate(cp['communities']['members'],1):rows.append([f'Komunitas {i}',str(len(m)),', '.join(m)])
t66=table_after(c66._p,rows)
c67=cap_after(t66._tbl,'Tabel 6.7 Tingkat keterlambatan komponen terhubung')
rows=[['Komponen','Pesanan','Tingkat keterlambatan','CI95']]
for c in cp['components']:
    if c['late_rate'] is None:continue
    rows.append([str(c['component_id']),n2(c['orders']),idn(c['late_rate'],4),f"[{idn(c['ci95'][0],4)}, {idn(c['ci95'][1],4)}]"])
t67=table_after(c67._p,rows)
c68=cap_after(t67._tbl,'Tabel 6.8 Tingkat keterlambatan menurut Days for shipment (scheduled) dan Shipping Mode')
rows=[['Nilai jadwal (hari)','Shipping Mode','Pesanan','Tingkat keterlambatan','CI95']]
mode_of={v['scheduled_value']:k for k,v in G['mode_rate'].items()}
for k,v in G['schedule_rate'].items():
    val=int(k.split('=')[1]);rows.append([str(val),mode_of.get(val,'-'),n2(v['n']),idn(v['late_rate'],4),f"[{idn(v['ci95'][0],4)}, {idn(v['ci95'][1],4)}]"])
t68=table_after(c68._p,rows)
pp=d.add_paragraph()
pp.add_run('Hubungan jadwal dan keterlambatan bersifat asosiatif. Nilai jadwal 1 hari (First Class) memiliki tingkat keterlambatan 1,0000, sedangkan 4 hari (Standard Class) 0,3985, sehingga jadwal berfungsi sebagai proksi mode pengiriman dan bukan penyebab langsung keterlambatan. Model dengan jadwal one-hot mencapai AP 0,821496 (CI95 [0,813022, 0,830012]). Varian ini disajikan sebagai eksperimen tambahan dan belum menggantikan desain utama.')
t68._tbl.addnext(pp._p)
# ====== 6.2: perbarui paragraf kontradiktif + insight ======
set_text(find_para('Relasi yang dianalisis menunjukkan kategori apa saja'),
 f"Relasi yang dianalisis menunjukkan kategori apa saja yang tercatat bersama dalam pesanan. Graf tidak memodelkan gudang, kurir, urutan perjalanan, atau waktu tempuh. Hubungan struktur graf dengan label diuji secara deskriptif melalui tingkat keterlambatan per kategori dan per komponen (Tabel 6.5 dan 6.7), bukan melalui pengujian kausal dan bukan sebagai fitur model. Tingkat keterlambatan per kategori berkisar {idn(cat_min,4)} (Category Id {cat_min_id}) sampai {idn(cat_max,4)} (Category Id {cat_max_id}), dengan keseluruhan pesanan sebesar {idn(overall,4)}. Rentang tersebut menunjukkan variasi antarkategori, tetapi tidak membuktikan bahwa kategori menyebabkan keterlambatan.")
set_text(find_para('Kategori dengan degree terbesar adalah'),
 f"Urutan sentralitas pada proyeksi kategori konsisten antarmetrik. Korelasi peringkat Kendall antara weighted degree dan PageRank adalah {idn(H['rank_similarity']['weighted_degree_vs_pagerank'],3)}, sedangkan antara weighted degree dan eigenvector {idn(H['rank_similarity']['weighted_degree_vs_eigenvector'],3)}. Kategori dengan degree terbesar adalah Category Id 17 dengan 19.534 pesanan unik, diikuti Category Id 18 dengan 17.954 pesanan dan Category Id 24 dengan 17.088 pesanan. Sebagai pembanding, betweenness eksak menempatkan Category Id 16 pada nilai tertinggi ({idn(H['betweenness_exact']['top10'][0]['betweenness'])}), berbeda dari urutan degree; antara lain karena antara merupakan ukuran jembatan lintasan, bukan jumlah koneksi. Nilai degree tidak menghitung kuantitas barang maupun jumlah kemunculan item kategori yang sama dalam satu pesanan.")
# Gambar 6.2 diganti (label + legenda)
tmp=d.add_paragraph();run=tmp.add_run();run.add_picture(str(OUT/'H_projection_fixed.png'),width=Inches(6.0));new_id=run._r.find('.//'+qn('a:blip')).get(qn('r:embed'));tmp._p.getparent().remove(tmp._p)
for i,el in enumerate(list(d.element.body)):
    if el.tag==qn('w:p') and el.findall('.//'+qn('w:drawing')):
        nxt=el.getnext()
        cap=''.join(t.text or '' for t in nxt.iter(qn('w:t'))) if nxt is not None and nxt.tag==qn('w:p') else ''
        if cap.startswith('Gambar 6.2'):
            for bl in el.findall('.//'+qn('a:blip')):bl.set(qn('r:embed'),new_id)
        if cap.startswith('Gambar 5.1'):
            tmp2=d.add_paragraph();r2=tmp2.add_run();r2.add_picture(str(OUT/'I_pr_roc_main.png'),width=Inches(6.2));nid2=r2._r.find('.//'+qn('a:blip')).get(qn('r:embed'));tmp2._p.getparent().remove(tmp2._p)
            for bl in el.findall('.//'+qn('a:blip')):bl.set(qn('r:embed'),nid2)
d.save(SRC)
print('5.6, Tabel 6.1, 6.4, 6.6-6.8 diperbarui; late rate range',round(cat_min,4),cat_min_id,'-',round(cat_max,4),cat_max_id,'overall',round(overall,4))
