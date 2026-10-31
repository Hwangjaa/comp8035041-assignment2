"""Regresi: cek kontradiksi hilang, angka lama vs baru, distribusi halaman."""
import json, re
from pathlib import Path
from docx import Document
import fitz
R=Path('/Users/rafaelyudianto/Documents/COMP8035041-Assignment-II')
O=R/'Assignment-II-DataCo-Laporan-Final-High-Quality.docx';PDF=R/'Assignment-II-DataCo-Laporan-Final-High-Quality.pdf'
d=Document(O);txt='\n'.join(p.text for p in d.paragraphs)
for t in d.tables:
    for r in t.rows:
        for c in r.cells:txt+='\n'+c.text
bad_phrases=['tidak dinyatakan ekuivalen','tidak menyetarakan kekuatan regularisasi','Tidak dilakukan pengujian hubungan struktur graf','tidak dilakukan uji signifikansi atau pengukuran ketidakpastian','tidak dilakukan pengukuran ketidakpastian','kontrak fitur','gerbang keputusan']
found=[b for b in bad_phrases if b.lower() in txt.lower()]
# angka utama baru harus ada
must_new=['0,736424','0,719782','0,307461','0,307467','2,0069843','0,868863','0,820887','0,821702','0,821496','0,4837','1,0000','0,3985','62','50,00','0,4963','0,6885','16','0,3161']
missing_new=[k for k in must_new if k not in txt]
# angka lama yang harus HILANG (konfigurasi Spark 0,1 sebagai threshold utama)
old_gone=['0.446602','0,446602','0.720601','0,720601','0.736851','0,736851']
still_old=[k for k in old_gone if k in txt]
rep={'contradictions_remaining':found,'missing_new_numbers':missing_new,'old_numbers_remaining':still_old}
# halaman sparse
pdf=fitz.open(PDF);sparse=[]
for i,pg in enumerate(pdf):
    t=pg.get_text().strip()
    if len(t)<300 and i>0:sparse.append((i+1,len(t)))
rep['pages']=len(pdf);rep['sparse_pages']=sparse
# tabel count + daftar
rep['tables']=len(d.tables);rep['toc_entries']=len([p for p in d.paragraphs if '\t' in p.text and p.text.rsplit('\t',1)[-1].isdigit()])
json.dump(rep,open(R/'analisis-lanjutan/REGRESSION2.json','w'),indent=2)
print(json.dumps(rep,indent=2,ensure_ascii=False))
print()
print('--- daftar tabel ---')
on=False
for p in d.paragraphs:
    t=p.text.strip()
    if t=='DAFTAR TABEL':on=True;continue
    if on and p.style.name.startswith('Heading'):break
    if on and '\t' in p.text:print('  ',p.text.replace('\t',' >> '))
