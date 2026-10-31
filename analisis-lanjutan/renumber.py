"""Renumber tabel evaluasi lama agar urutan fisik = numerik (5.4, 5.5 -> 5.5, 5.6)."""
import re
from docx import Document
from docx.oxml.ns import qn
from pathlib import Path
R=Path('/Users/rafaelyudianto/Documents/COMP8035041-Assignment-II')
O=R/'Assignment-II-DataCo-Laporan-Final-High-Quality.docx'
d=Document(O)
def rename_smart(old,new):
    cnt=0
    for p in d.paragraphs:
        if old not in p.text:continue
        prefix=[]
        for r in p.runs:
            if any(x.tag not in (qn('w:rPr'),qn('w:t')) for x in r._r):break
            prefix.append(r)
        joined=''.join(r.text for r in prefix)
        if old in joined:
            prefix[0].text=joined.replace(old,new)
            for r in prefix[1:]:r.text=''
            cnt+=1
    return cnt
a=rename_smart('Tabel 5.5 Hasil evaluasi pada data validation','Tabel 5.6 Hasil evaluasi pada data validation')
b=rename_smart('Tabel 5.4 Hasil evaluasi pada data test','Tabel 5.5 Hasil evaluasi pada data test')
# referensi dalam teks
c=0
for p in d.paragraphs:
    if 'menyajikan metrik' in p.text and 'Tabel 5.4' in p.text:
        prefix=[]
        for r in p.runs:
            if any(x.tag not in (qn('w:rPr'),qn('w:t')) for x in r._r):break
            prefix.append(r)
        if prefix:
            j=''.join(r.text for r in prefix)
            prefix[0].text=re.sub(r'Tabel 5\.4(\s*[\u2013-]\s*)5\.5',lambda m:'Tabel 5.5'+m.group(1)+'5.6',j)
            for r in prefix[1:]:r.text=''
            c+=1
d.save(O)
print('renamed validation->5.6:',a,'test->5.5:',b,'ref:',c)
# verifikasi
d2=Document(O);on=False
for p in d2.paragraphs:
    t=p.text.strip()
    if t=='DAFTAR TABEL':on=True;continue
    if t=='DAFTAR GAMBAR':break
    if on and '\t' in p.text:print('  ',p.text.replace('\t',' >> '))
