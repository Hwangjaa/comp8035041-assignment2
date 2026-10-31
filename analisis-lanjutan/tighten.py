"""Rapikan tata letak: rapatkan entri daftar, cegah tabel terpotong & caption terpisah."""
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from pathlib import Path
R=Path('/Users/rafaelyudianto/Documents/COMP8035041-Assignment-II')
O=R/'Assignment-II-DataCo-Laporan-Final-High-Quality.docx'
d=Document(O)
# 1. Rapatkan semua entri daftar (DAFTAR ISI/TABEL/GAMBAR) agar tidak meluber
start=False;n=0
for p in d.paragraphs:
    t=p.text.strip()
    if t in ('DAFTAR ISI','DAFTAR TABEL','DAFTAR GAMBAR'):start=True;continue
    if p.style.name.startswith('Heading') and t.startswith('BAB 1.'):start=False
    if start and '\t' in p.text and p.text.rsplit('\t',1)[-1].isdigit():
        pf=p.paragraph_format
        pf.space_before=Pt(0);pf.space_after=Pt(0.5);pf.line_spacing=0.92
        for r in p.runs:r.font.size=Pt(10.5)
        n+=1
# 2. Caption tabel menempel ke tabel; baris tidak terpotong
def set_cantsplit(tbl):
    for row in tbl.rows:
        trPr=row._tr.get_or_add_trPr()
        if trPr.find(qn('w:cantSplit')) is None:trPr.append(OxmlElement('w:cantSplit'))
caps=0
for p in d.paragraphs:
    if p.style.name=='Table Caption':
        p.paragraph_format.keep_with_next=True;caps+=1
for tbl in d.tables:
    set_cantsplit(tbl)
    if len(tbl.rows)<=12:
        for row in tbl.rows[:-1]:
            for c in row.cells:
                for pp in c.paragraphs:pp.paragraph_format.keep_with_next=True
d.save(O)
print('toc entries tightened:',n,'table captions keep_with_next:',caps,'tables:',len(d.tables))
