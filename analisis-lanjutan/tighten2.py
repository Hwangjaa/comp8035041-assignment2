"""Paksa DAFTAR TABEL & DAFTAR GAMBAR muat satu halaman (spasi sangat rapat)."""
from docx import Document
from docx.shared import Pt
from pathlib import Path
R=Path('/Users/rafaelyudianto/Documents/COMP8035041-Assignment-II')
O=R/'Assignment-II-DataCo-Laporan-Final-High-Quality.docx'
d=Document(O)
def tighten(header,stop):
    on=False;n=0
    for p in d.paragraphs:
        t=p.text.strip()
        if t==header:on=True
        elif on and p.style.name.startswith('Heading') and t!=header:on=False
        if on and '\t' in p.text and p.text.rsplit('\t',1)[-1].isdigit():
            pf=p.paragraph_format
            pf.space_before=Pt(0);pf.space_after=Pt(0);pf.line_spacing=0.85
            for r in p.runs:r.font.size=Pt(10)
            n+=1
    return n
a=tighten('DAFTAR TABEL','DAFTAR GAMBAR');b=tighten('DAFTAR GAMBAR','BAB 1.')
# heading daftar: kurangi space_after
for p in d.paragraphs:
    if p.text.strip() in ('DAFTAR TABEL','DAFTAR GAMBAR','DAFTAR ISI'):
        p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(0)
d.save(O)
print('tabel entries:',a,'gambar entries:',b)
