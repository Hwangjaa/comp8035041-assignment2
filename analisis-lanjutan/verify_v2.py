from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
import fitz, re, json, hashlib
R=Path('/Users/rafaelyudianto/Documents/COMP8035041-Assignment-II')
O=R/'Assignment-II-DataCo-Laporan-Final-High-Quality.docx'
PDF=R/'Assignment-II-DataCo-Laporan-Final-High-Quality.pdf'
b=Document(O);pdf=fitz.open(PDF)
norm=lambda x:re.sub(r'\s+','',x).casefold();pt=[norm(x.get_text()) for x in pdf]
entries=[];bad=[]
for p in b.paragraphs:
    if '\t' in p.text and p.text.rsplit('\t',1)[-1].isdigit():
        t,n=p.text.rsplit('\t',1);n=int(n);entries.append((t,n))
        if not(0<n<=len(pt)) or norm(t) not in pt[n-1]:bad.append((t,n))
empty=[i+1 for i,pg in enumerate(pdf) if len(pg.get_text().strip())<10]
geo=[]
for i,pg in enumerate(pdf):
    for bl in pg.get_text('dict')['blocks']:
        x0,y0,x1,y1=bl['bbox']
        if min(x0,y0)<-0.5 or x1>pg.rect.width+0.5 or y1>pg.rect.height+0.5:geo.append(i+1)
# urutan: halaman harus non-turun dalam tiap daftar
def list_pages(start_marker,end_marker):
    ps=[];on=False
    for p in b.paragraphs:
        if p.text.strip()==start_marker:on=True;continue
        if p.text.strip()==end_marker:on=False
        if on and '\t' in p.text and p.text.rsplit('\t',1)[-1].isdigit():ps.append(int(p.text.rsplit('\t',1)[-1]))
    return ps
tbls=list_pages('DAFTAR TABEL','DAFTAR GAMBAR');figs=list_pages('DAFTAR GAMBAR','BAB 1. PENDAHULUAN')
isi=[n for t,n in entries if not t.startswith(('Tabel ','Gambar '))]
txt='\n'.join(p.text for p in b.paragraphs)
rep={'pages':len(pdf),'sha256':hashlib.sha256(O.read_bytes()).hexdigest(),
 'toc_entries':len(entries),'toc_bad':bad[:8],
 'daftar_tabel_sorted':tbls==sorted(tbls),'daftar_gambar_sorted':figs==sorted(figs),
 'isi_sorted':isi==sorted(isi),
 'empty_pages':empty,'geometry_pages':sorted(set(geo)),'tables':len(b.tables),'images':len(b.inline_shapes),
 'internal_codes':[c for c in re.findall(r'\b(?:D2|H1|M0|F[2-8]|R0|T0|U0|G1)\b',txt)][:5],
 'placeholder':[w for w in ['kontrak fitur','TODO','placeholder','gerbang keputusan','keputusan mahasiswa'] if w.lower() in txt.lower()],
 'under_40':len(pdf)<=40}
json.dump(rep,open(R/'analisis-lanjutan/FINAL-VERIFY.json','w'),indent=2)
print(json.dumps(rep,indent=2,ensure_ascii=False))
