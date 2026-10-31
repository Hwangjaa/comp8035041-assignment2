from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
import fitz, re, hashlib, shutil
R=Path('/Users/rafaelyudianto/Documents/COMP8035041-Assignment-II')
O=R/'Assignment-II-DataCo-Laporan-Final-High-Quality.docx'
PDF=Path('/Users/rafaelyudianto/.hermes/cache/scratch/dataco-hq/Assignment-II-DataCo-Laporan-Final-High-Quality.pdf')
pdf=fitz.open(PDF);norm=lambda x:re.sub(r'\s+','',x).casefold()
pages=[norm(p.get_text()) for p in pdf]
TOC=set(range(0,4))
b=Document(O);fixed=0;unresolved=[]
for p in b.paragraphs:
    t=p.text
    if '\t' in t and t.rsplit('\t',1)[-1].isdigit():
        title=t.rsplit('\t',1)[0];key=norm(title);target=None
        for i,tx in enumerate(pages):
            if i in TOC:continue
            if key in tx:target=i+1;break
        ts=[x for x in p._p.iter(qn('w:t'))]
        if target and ts and ts[-1].text is not None:
            ts[-1].text=str(target);fixed+=1
        else:unresolved.append((title,target))
b.save(O)
print('cache updated:',fixed,'unresolved:',unresolved[:5])
# copy Word-rendered PDF as deliverable PDF
dst=R/'Assignment-II-DataCo-Laporan-Final-High-Quality.pdf'
shutil.copy2(PDF,dst)
print('pdf copied. docx sha256:',hashlib.sha256(O.read_bytes()).hexdigest())
d2=Document(O)
v=[int(p.text.rsplit('\t',1)[-1]) for p in d2.paragraphs if '\t' in p.text and p.text.rsplit('\t',1)[-1].isdigit()]
print('entries:',len(v),'distinct:',sorted(set(v)))
