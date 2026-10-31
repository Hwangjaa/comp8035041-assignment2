from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
import fitz, re, hashlib, shutil
R=Path('/Users/rafaelyudianto/Documents/COMP8035041-Assignment-II')
# mulai dari docx yang sudah punya tabel/gambar/teks baru (punya cache buruk, akan ditimpa)
O=R/'Assignment-II-DataCo-Laporan-Final-High-Quality.docx'
PDF=Path('/Users/rafaelyudianto/.hermes/cache/scratch/dataco-hq/Assignment-II-DataCo-Laporan-Final-High-Quality.pdf')
pdf=fitz.open(PDF);norm=lambda x:re.sub(r'\s+','',x).casefold()
pages=[norm(p.get_text()) for p in pdf]
b=Document(O);fixed=0;unres=[]
for p in b.paragraphs:
    t=p.text
    if '\t' in t and t.rsplit('\t',1)[-1].isdigit():
        title=t.rsplit('\t',1)[0];key=norm(title)
        hits=[i+1 for i,tx in enumerate(pages) if key in tx]
        if not hits:unres.append(title);continue
        target=hits[-1]  # kemunculan terakhir = badan
        ts=[x for x in p._p.iter(qn('w:t'))]
        if ts and ts[-1].text is not None:ts[-1].text=str(target);fixed+=1
        else:unres.append(title)
b.save(O)
print('fixed:',fixed,'unresolved:',unres)
d2=Document(O)
v=[int(p.text.rsplit('\t',1)[-1]) for p in d2.paragraphs if '\t' in p.text and p.text.rsplit('\t',1)[-1].isdigit()]
print('entries:',len(v),'distinct:',sorted(set(v)))
print('sha256:',hashlib.sha256(O.read_bytes()).hexdigest())
