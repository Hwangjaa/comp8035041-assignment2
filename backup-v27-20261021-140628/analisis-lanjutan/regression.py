"""Fase 4: regresi konsistensi angka. Cari angka lama/usang di DOCX final."""
from pathlib import Path
from docx import Document
import re, json
R=Path('/Users/rafaelyudianto/Documents/COMP8035041-Assignment-II')
O=R/'Assignment-II-DataCo-Laporan-Final-High-Quality.docx'
d=Document(O)
txt='\n'.join(p.text for p in d.paragraphs)
for t in d.tables:
    for r in t.rows:
        for c in r.cells:txt+='\n'+c.text
# angka acuan yang HARUS muncul
must=['180.519','65.752','172.765','62.897','36.048','26.849','7.754','2.855','49.826','5.028','8.043','48','62.948','152.676','469.977','466.728','3.249']
present={k:(k in txt) for k in must}
# angka metrik lama harus konsisten (bukan hilang)
old_metrics=['0.736424','0.736851','0.720601','0.719782','0.576402','0.731288','3.407','4.636']
old_present={k:(k in txt) for k in old_metrics}
# placeholder / editorial leaks
bad_phrases=['TODO','TBD','XXX','lorem','placeholder','[isi','catatan internal','keputusan mahasiswa','kontrak fitur','gerbang keputusan']
found_bad=[p for p in bad_phrases if p.lower() in txt.lower()]
rep={'reference_numbers_present':present,'old_metrics_present':old_present,'placeholder_leaks':found_bad,
     'missing_reference':[k for k,v in present.items() if not v]}
json.dump(rep,open(R/'analisis-lanjutan/REGRESSION.json','w'),indent=2)
print(json.dumps(rep,indent=2,ensure_ascii=False))
