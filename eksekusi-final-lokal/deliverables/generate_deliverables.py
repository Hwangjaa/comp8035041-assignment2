from pathlib import Path
import json, hashlib, re, zipfile, copy
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pptx import Presentation
from pptx.util import Inches as PI, Pt as PP
from pptx.dml.color import RGBColor
ROOT=Path('/Users/rafaelyudianto/Documents/COMP8035041-Assignment-II')
OUT=ROOT/'eksekusi-final-lokal/deliverables'; RUN=ROOT/'eksekusi-final-lokal/run-actual'
SRC=ROOT/'laporan-draft/Assignment II - DataCo - Draf Laporan Lengkap.docx'
load=lambda n:json.loads((RUN/n).read_text())
metrics=load('metrics.json'); fc=load('feature-contract.json'); graph=load('graph-summary.json'); status=load('status.json'); env=load('environment.json'); split=pd.read_csv(RUN/'split.csv'); degree=pd.read_csv(RUN/'graph-category-degree.csv'); components=pd.read_csv(RUN/'graph-components.csv')
assert status['exit_status']==0 and fc['n_features']==48 and len(fc['source_columns'])==5 and len(fc['feature_names'])==48
assert int(split.total.sum())==62897 and int(degree.degree.sum())==graph['edges'] and int(components.nodes.sum())==graph['nodes']
sha=lambda p:hashlib.sha256(Path(p).read_bytes()).hexdigest()
before=sha(SRC)
doc=Document(SRC); ps=list(doc.paragraphs)
def settext(p,text):
    assert not p._p.xpath('.//w:drawing|.//m:oMath|.//w:fldChar|.//w:instrText'), 'unsafe replacement'
    # Keep paragraph properties and source run styling.
    rp=copy.deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
    for c in list(p._p):
        if c.tag!=qn('w:pPr'): p._p.remove(c)
    r=p.add_run(str(text))
    if rp is not None:r._r.insert(0,rp)
updates={
1:'Studi Kasus: DataCo Smart Supply Chain — Assignment II (Hasil Eksekusi Lokal)',
25:'Proyek ini melanjutkan DataCo Smart Supply Chain dari Assignment I untuk pelaporan risiko keterlambatan pengiriman. Pembersihan, pembentukan fitur, pelatihan Logistic Regression pada scikit-learn dan Spark MLlib, evaluasi serta graph telah dieksekusi secara lokal. Angka aktual dilaporkan pada Bab 5–6; interpretasi dan kesimpulan eksperimen wajib ditulis mahasiswa.',
28:'Akturk et al. (2022) melaporkan hubungan durasi pemrosesan dan pengiriman dengan penilaian pelanggan. Literatur tersebut bukan pengukuran dampak bisnis DataCo. Proyek menghasilkan probabilitas dan prediksi historis, tetapi belum mengukur biaya, kepuasan pelanggan, keberhasilan intervensi atau penggunaan daftar prioritas operasional.',
51:'Bukti persiapan Colab sebelumnya terpisah dari eksekusi final. Eksekusi final berlangsung lokal pada macOS, Python 3.12.15, Spark/PySpark 3.5.9, Java 17.0.20.1, mode local[2], dengan fs.defaultFS=file:///. Lokasi berkas lokal bukan HDFS dan keberadaan pustaka Hadoop di Spark bukan bukti layanan HDFS aktif. Versi lengkap tersimpan pada environment.json dan spark-environment.json.',
53:'Gambar 2.1 dan Gambar 2.2 dipertahankan sebagai arsitektur konseptual dan bukti historis Assignment I, bukan diagram implementasi final lengkap. Eksekusi final Assignment II menggunakan berkas lokal file:///, bukan HDFS pada diagram. Model, evaluasi dan graph telah dieksekusi; penggunaan operasional dan streaming belum diuji.',
58:'Pada jalur historis, pembersihan meliputi pemetaan nama, parsing tanggal, penghapusan empat atribut, late_days untuk audit dan cakupan nonbatal. Eksekusi final membentuk satu observasi per Order Id, tepat lima kelompok fitur, probabilitas serta prediksi tersimpan. Agregasi prioritas operasional dan streaming tidak dijalankan; CSV historis bukan bukti layanan streaming.',
91:'Data mentah dan keluaran audit Assignment I tetap terpisah. Eksekusi final menyimpan orders.csv, split.csv, feature-contract.json, model, prediksi, metrik dan graph dalam eksekusi-final-lokal/run-actual. CSV audit tingkat item bukan matriks fitur. Bukti HDFS historis hanya /smoke/input.txt; dataset lengkap belum dimuat ke HDFS.',
94:'Assignment II menggunakan satu observasi per Order Id. Eksekusi lokal memeriksa konsistensi target, jadwal, segmen, market dan tanggal di dalam pesanan; konflik singleton dan target berjumlah 0. Agregasi kategori memakai himpunan ID unik dan jumlah item menghitung Order Item Id unik. Metode tersebut melaksanakan keputusan mahasiswa, bukan memilih metode baru.',
101:'Durasi aktual, shipping_date, late_days, Delivery Status dan status akhir tidak digunakan sebagai prediktor. Label hanya target; Order Id hanya kunci. Kontrak aktual mencantumkan tepat lima kelompok fitur dan vocabulary dipelajari hanya pada training. Hal ini tidak membuktikan ketersediaan nilai as-of pada waktu prediksi operasional. Split temporal serta metrik aktual dilaporkan pada Bab 5; pembahasan leakage mengacu Kaufman et al. (2012) dan Kapoor dan Narayanan (2023).',
103:'Bagian 4.3 dan Bab 5–6 menyajikan hasil eksekusi lokal yang tercatat dalam run-actual. Pelaporan angka dipisahkan dari justifikasi metode, interpretasi, perbandingan substantif dan kesimpulan yang menjadi pekerjaan mahasiswa.',
107:'Audit aktual menghasilkan 62.897 pesanan nonbatal dari 172.765 item; 2.855 pesanan dan 7.754 item dikeluarkan. Pesanan campuran batal/nonbatal berjumlah 0. Target tingkat pesanan adalah 36.048 label 1 dan 26.849 label 0. Satu pesanan adalah satu observasi; jumlah item bukan jumlah observasi model.',
112:'Jejak pembersihan Assignment I dipertahankan. Eksekusi final membaca CSV sumber, mengecualikan Shipping canceled, memeriksa konflik singleton dan label, mengagregasi per Order Id serta membuat split temporal. orders.csv dan audit-orders.json menjadi keluaran audit aktual. Data audit tingkat item tidak langsung diperlakukan sebagai matriks fitur.',
113:'Eksekusi lokal memeriksa SHA-256 ketiga CSV terhadap manifest lokal sebelumnya, bukan mengambil ulang metadata penerbit. Terdapat 180.519 item sumber dan 65.752 pesanan sumber; konflik singleton, konflik target serta kategori kosong pada cakupan dipertahankan berjumlah 0. Rincian environment dan manifest tersedia dalam run-actual.',
114:'Vocabulary dipelajari hanya pada training. MISSING dipakai untuk kategori kosong dan UNKNOWN untuk nilai di luar vocabulary training. Tidak digunakan scaling atau transformasi log; dua fitur numerik memakai satuan asli. Nilai numerik atau identitas wajib yang kosong/nonfinite menghentikan pipeline. Transformasi yang sama diterapkan ke validation dan test.',
116:'Kontrak aktual memuat tepat lima kelompok fitur: Days for shipment (scheduled), Customer Segment, Market, Category Id multi-hot, serta jumlah Order Item Id unik per pesanan. Jadwal, segmen dan market diambil sebagai nilai tunggal setelah pemeriksaan konsistensi. Category Id direpresentasikan sebagai himpunan kategori per pesanan. Quantity, atribut keuangan dan turunannya tidak diaktifkan.',
119:'Encoding aktual menghasilkan 48 dimensi: 2 numerik, 5 dimensi Customer Segment, 7 dimensi Market dan 34 dimensi Category Id, termasuk MISSING dan UNKNOWN. Ukuran matriks adalah training 49.826 × 48, validation 5.028 × 48 dan test 8.043 × 48. [WAJIB MAHASISWA: jelaskan justifikasi fitur dan agregasi serta batas ketersediaan as-of; jangan menyatakan terbukti operasional.]',
123:'Parameter dilaksanakan sesuai keputusan mahasiswa tanpa tuning. C=1.0 pada scikit-learn dan regParam=0.1 pada Spark tidak dinyatakan ekuivalen. [WAJIB MAHASISWA: tulis justifikasi algoritma dan parameter berdasarkan pemahaman sendiri.]',
127:'Tabel 5.3 berasal dari split.csv hasil eksekusi final. orders.csv memiliki kunci Order Id unik. Kedua jalur menerima urutan fitur yang sama; probabilitas Spark dipetakan kembali melalui Order Id. Model dipasang hanya pada training dan tidak di-refit menggunakan gabungan training–validation (R0).',
128:'Lingkungan aktual lokal: Python 3.12.15, pandas 3.0.6, NumPy 2.5.4, SciPy 1.18.1, scikit-learn 1.9.1, PySpark/Spark 3.5.9, NetworkX 3.7, Java 17.0.20.1. Spark berjalan local[2] dengan file:///. Catatan persiapan Colab tidak digunakan sebagai versi eksekusi final; requirements-observed.txt merupakan inventaris teramati, bukan jaminan lockfile reproduksi lintas mesin.',
129:'status.json mencatat exit_status=0 dan training_evaluation_graph=executed. Model scikit-learn tercatat 415 iterasi dan Spark 71 iterasi pada convergence.json. Model, parameter, probabilitas, prediksi dan metrik tersimpan dalam run-actual. B0 berarti benchmark tidak dilakukan; waktu dalam log bukan pengukuran speedup atau skalabilitas multi-node.',
132:'H1 memilih threshold dengan F1 validation maksimum dari probabilitas unik ditambah kandidat all-negative; ties memilih threshold terbesar dan kelas positif memakai probabilitas >= threshold. Threshold dibekukan sebelum prediksi test: baseline 0.5724722032673705; scikit-learn 0.3074755313277422; Spark 0.4466021438573422. U0 bersifat deskriptif tanpa uji signifikansi.',
133:'Tabel 5.4 Hasil test aktual; threshold dipilih hanya pada validation',
134:'Sumber angka Tabel 5.4 adalah metrics.json; angka metrik ditampilkan enam desimal, bukan presisi penuh artefak. Setiap CSV predictions-{validation|test}-{baseline|sklearn_lr|spark_lr}.csv menyimpan Order Id, tanggal, partisi, target, model, probability_1, threshold dan prediction. Parameter internal Spark threshold=0.5 bukan threshold evaluasi H1: evaluasi menggunakan probability_1 dan threshold yang dibekukan.',
136:'[WAJIB DITULIS MAHASISWA: interpretasi hasil aktual, perbandingan kinerja baseline dan kedua implementasi, makna FP/FN, konsekuensi threshold, argumentasi serta kesimpulan. AI hanya menyalin keluaran faktual pada Bagian 5.5; bagian analisis ini belum selesai dan bukan kesimpulan eksperimen.]',
141:'Eksekusi graph penuh menghasilkan 62.948 node: 62.897 pesanan dan 51 kategori; 152.676 edge unik; 19 connected components; self-loop dan duplicate edge masing-masing 0. Pembentukan edge memakai Spark DataFrame distinct; degree dan komponen dihitung dengan NetworkX pada satu mesin. Identitas node kategori tetap Category Id, bukan penggabungan berdasarkan nama.',
143:'Degree kategori menghitung tetangga pesanan unik, bukan kuantitas barang. Subset visual berisi 10 kategori, 200 pesanan, 210 node dan 495 edge; spring layout seed=42 dan iterations=50. Statistik graph penuh dihitung sebelum pemilihan subset. graph-edges.csv, graph-category-degree.csv, graph-components.csv dan graph-component-membership.csv menyimpan hasil penuh.',
145:'[WAJIB DITULIS MAHASISWA: analisis dan insight graph berdasarkan keluaran aktual pada Bagian 6.1, alasan hubungan yang relevan, batas representasi dan kesimpulan. Angka degree serta komponen yang disalin bukan interpretasi, pembuktian sebab-akibat atau rekomendasi prioritas.]',
146:'Graph merepresentasikan keanggotaan pesanan–kategori, bukan rute fisik pengiriman. Relasi ini tidak membuktikan penyebab keterlambatan atau dampak intervensi. Graph tidak dimasukkan sebagai fitur supervised; interpretasi graph dan evaluasi model tetap pekerjaan mahasiswa.',
156:'Dataset: https://doi.org/10.17632/8gx2fvg2k6.5. Notebook tereksekusi lokal: eksekusi-final-lokal/Assignment-II-DataCo-executed-local.ipynb; kode: eksekusi-final-lokal/src/pipeline.py; hasil: eksekusi-final-lokal/run-actual. Slide: deliverables/Assignment-II-DataCo-Hasil-Aktual.pptx; naskah: deliverables/Naskah-Presentasi-Faktual.md. [WAJIB PENULIS: tautan berbagi dataset/kode/slide yang dapat diakses penilai dan tautan video Assignment II setelah direkam; video teknis ditangani terpisah; draf teknis AI bukan final mahasiswa.]',
167:'Bantuan AI meliputi penyuntingan, penataan salinan laporan, pelaporan keluaran aktual, pembuatan slide dan naskah faktual. Riwayat diskusi, bantuan kode, visualisasi dan prompt wajib ditinjau serta diungkapkan sesuai penggunaan sebenarnya. Analisis, justifikasi, perbandingan dan kesimpulan wajib ditulis mahasiswa. Persentase kontribusi akhir: [ditetapkan mahasiswa; batas brief maksimal 50%, bukan persentase otomatis].'
}
for i,t in updates.items():settext(ps[i],t)
def cell(t,r,c,text):
    ce=doc.tables[t].cell(r,c); settext(ce.paragraphs[0],text)
    for p in ce.paragraphs[1:]:settext(p,'')
for r,c,t in [(3,1,'Pembersihan Python historis; pembentukan edge graph Spark; fitur, fit dan evaluasi lokal selesai'),(3,2,'Streaming dan benchmark multi-node belum dilakukan'),(4,1,'Probabilitas, prediksi, metrik dan graph tersedia dalam run-actual'),(4,2,'Penggunaan operasional, intervensi dan analisis mahasiswa belum selesai')]:cell(2,r,c,t)
cell(8,3,2,'orders.csv, split.csv, feature-contract.json, model dan prediksi tersedia lokal')
cell(9,1,1,'Tepat lima kelompok: jadwal, segmen, market, kategori multi-hot, jumlah ID item unik');cell(9,1,2,'Kontrak aktual; vocabulary training-only; ketersediaan as-of operasional belum terbukti')
cell(10,4,3,'Model langsung per pesanan telah dilatih dan dievaluasi; penggunaan prioritas operasional belum diuji')
cell(11,6,1,'Satu observasi per Order Id; 48 dimensi aktual')
feature_rows=[['Komponen','Kontrak aktual'],['Unit dan target','Order Id; Late_delivery_risk'],['Lima kelompok fitur','Jadwal; segmen; market; Category Id multi-hot; jumlah Order Item Id unik'],['Quantity dan keuangan','Tidak diaktifkan, termasuk turunannya'],['Statistik transformasi','Vocabulary hanya fit pada training; MISSING dan UNKNOWN'],['Matriks input final','49.826 × 48; 5.028 × 48; 8.043 × 48; tanpa scaling/log']]
for r,row in enumerate(feature_rows):
 for c,v in enumerate(row):cell(12,r,c,v)
for r,row in enumerate(split.itertuples(index=False),1):
 for c,v in enumerate([{'train':'Training','validation':'Validation','test':'Test'}[row.partition],row.label_0,row.label_1,row.total]):cell(14,r,c,v)
mt={m['model']:m for m in metrics if m['partition']=='test'}
keys=['average_precision','roc_auc','precision','recall','binary_f1','confusion_matrix','threshold']
for r,key in enumerate(keys,1):
 for c,name in enumerate(['baseline','sklearn_lr','spark_lr'],1):
  v=mt[name][key];cell(15,r,c,', '.join(map(str,v)) if isinstance(v,list) else f'{v:.6f}')
cell(17,3,2,'Diagram historis dipertahankan; graph-subset.png aktual tersedia; verifikasi pengungkapan penulis')
cell(17,2,2,'Laporan keluaran aktual, penataan, slide dan naskah faktual; bukan analisis akademik.')
# Add factual validation and graph inventories near the relevant sections.
def insert_table_before(anchor,headers,rows):
 caption=anchor.insert_paragraph_before({'Validation':'Tabel 5.5 Hasil validation aktual','Category Id':'Tabel 6.2 Degree sepuluh kategori pada keluaran aktual','Komponen':'Tabel 6.3 Komponen graph penuh aktual'}[headers[0]])
 caption.style='Table Caption'
 caption.paragraph_format.keep_with_next=True
 t=doc.add_table(rows=1, cols=len(headers)); t.style='Table Grid'
 for c,v in enumerate(headers):t.rows[0].cells[c].text=str(v)
 for row in rows:
  for c,v in enumerate(row):t.add_row() if False else None
  cells=t.add_row().cells
  for c,v in enumerate(row):cells[c].text=str(v)
 anchor._p.addprevious(t._tbl)
 return t
insert_table_before(ps[135],['Validation','AP','ROC-AUC','Precision','Recall','F1','TN, FP, FN, TP'],[[m['model']]+[f'{m[k]:.6f}' for k in keys[:5]]+[str(m['confusion_matrix'])] for m in metrics if m['partition']=='validation'])
insert_table_before(ps[144],['Category Id','Degree (pesanan unik)'],degree.head(10).values.tolist())
insert_table_before(ps[144],['Komponen','Node','Pesanan','Kategori','Edge'],components.values.tolist())
# Tables and chapter pagination rules; preserve existing pictures, fields, native math.
for p in doc.paragraphs:
 if p.style.name.startswith('Heading'):
  p.paragraph_format.keep_with_next=True
  if p.text.startswith('BAB '):p.paragraph_format.page_break_before=True
for t in doc.tables:
 for row in t.rows:
  pr=row._tr.get_or_add_trPr()
  if pr.find(qn('w:cantSplit')) is None:pr.append(OxmlElement('w:cantSplit'))
# Existing captions are styled text rather than SEQ fields: list by actual styles.
for ins in doc.element.xpath('//w:instrText'):
 if 'TOC' in ins.text and ('Tabel' in ins.text or 'Table' in ins.text):ins.text=' TOC \\h \\z \\t "Table Caption,1" '
 elif 'TOC' in ins.text and ('Gambar' in ins.text or 'Figure' in ins.text):ins.text=' TOC \\h \\z \\t "Figure Caption,1" '
# Instruct Word to calculate fields; no fabricated TOC page numbers.
u=OxmlElement('w:updateFields');u.set(qn('w:val'),'true');doc.settings.element.append(u)
# Controlled foreign terms; keep literal identifiers and all special XML untouched.
terms=['machine learning','graph analytics','Logistic Regression','connected components','self-loop','duplicate edge','training','validation','test','threshold','scaling','multi-hot','vocabulary','baseline','benchmark','streaming','snapshot','pipeline','deployment','class weighting','resampling','feature engineering']
pat=re.compile(r'(?<![\w:/.-])('+'|'.join(re.escape(t) for t in sorted(terms,key=len,reverse=True))+r')(?![\w/.-])',re.I)
for p in list(doc.paragraphs)+[p for t in doc.tables for row in t.rows for c in row.cells for p in c.paragraphs]:
 for r in list(p.runs):
  if r._r.xpath('.//w:drawing|.//m:oMath|.//w:fldChar|.//w:instrText'):continue
  text=r.text
  if not pat.search(text):continue
  pos=0
  for m in pat.finditer(text):
   for seg,italic in [(text[pos:m.start()],False),(m.group(),True)]:
    if seg:
     nr=copy.deepcopy(r._r);nr.xpath('./w:t')[0].text=seg;nr.xpath('./w:t')[0].set(qn('xml:space'),'preserve')
     for extra in nr.xpath('./w:t')[1:]:nr.remove(extra)
     if italic:
      rp=nr.get_or_add_rPr();it=OxmlElement('w:i');rp.append(it)
     r._r.addprevious(nr)
   pos=m.end()
  if text[pos:]:
   nr=copy.deepcopy(r._r);nr.xpath('./w:t')[0].text=text[pos:];nr.xpath('./w:t')[0].set(qn('xml:space'),'preserve');r._r.addprevious(nr)
  r._r.getparent().remove(r._r)
report=OUT/'Assignment-II-DataCo-Laporan-Hasil-Aktual.docx';doc.save(report)
# Factual deck and synchronized speaker script: 8m45 including student-owned segment.
slides=[
('Assignment II | Hasil aktual lokal',['Rafael Yudianto | 2702752495 | COMP8035041','DataCo Smart Supply Chain v5','Laporan angka eksekusi; analisis dan kesimpulan diisi mahasiswa.'],'00:00–00:35','Presentasi ini melaporkan keluaran aktual eksekusi lokal Assignment II. Data menggunakan DataCo Smart Supply Chain versi 5. Saya membedakan hasil yang tercatat dalam artefak dari interpretasi serta kesimpulan yang harus saya jelaskan sendiri.'),
('Cakupan dan unit observasi',['180.519 item / 65.752 pesanan sumber','172.765 item / 62.897 pesanan nonbatal dipertahankan','2.855 pesanan dikeluarkan; pesanan campuran = 0','Label pesanan: 0 = 26.849; 1 = 36.048'],'00:35–01:15','Sumber transaksi memuat 180.519 item dan 65.752 pesanan. Setelah cakupan Shipping canceled dikeluarkan, tersisa 172.765 item dalam 62.897 pesanan. Terdapat 2.855 pesanan yang dikeluarkan dan tidak ada pesanan campuran batal serta nonbatal. Unit model adalah satu pesanan, bukan satu item. Target sumber dipertahankan; istilah dispatch tidak membuktikan penyerahan fisik kepada kurir.'),
('Lima kelompok fitur → 48 dimensi',['Days for shipment (scheduled) dan jumlah Order Item Id unik','Customer Segment; Market; Category Id multi-hot','Vocabulary hanya training; MISSING / UNKNOWN','48 = 2 numerik + 5 segmen + 7 market + 34 kategori','Tanpa quantity, keuangan, scaling atau log'],'01:15–02:00','Eksekusi memakai tepat lima kelompok fitur sesuai kontrak mahasiswa. Dua kelompok numerik adalah durasi terjadwal dan jumlah Order Item Id unik. Tiga kelompok kategoris adalah segmen pelanggan, market dan himpunan Category Id. Vocabulary dipelajari hanya pada training. MISSING dan UNKNOWN tersedia. Kontrak aktual mencatat 48 dimensi. Ini bukan pembuktian ketersediaan fitur as-of pada sistem operasional.'),
('Split dan pelatihan',['Train < 2017-07-01: 49.826 × 48','Validation 2017-07-01 s.d. < 2017-10-01: 5.028 × 48','Test ≥ 2017-10-01: 8.043 × 48','LR scikit-learn C=1.0; LR Spark regParam=0.1','Tanpa tuning / refit; parameter tidak ekuivalen'],'02:00–02:45','Split mengikuti order_date. Training sebelum Juli 2017, validation dari Juli sampai sebelum Oktober 2017, dan test mulai Oktober 2017. Matriks pada ketiganya memiliki 48 kolom. Model menggunakan Logistic Regression pada scikit-learn dan Spark MLlib. C bernilai 1 dan regParam bernilai 0,1 tidak dinyatakan ekuivalen. Tidak dilakukan tuning atau refit pada gabungan training dan validation. Baseline menyimpan prior positif training.'),
('Evaluasi test aktual',['Model | AP | ROC-AUC | F1']+[f"{n}: {mt[n]['average_precision']:.6f} | {mt[n]['roc_auc']:.6f} | {mt[n]['binary_f1']:.6f}" for n in mt]+['Semua: precision=0.576402; recall=1.000000','Semua: TN=0; FP=3.407; FN=0; TP=4.636'],'02:45–03:40','Slide ini menyalin angka test tanpa memberi penilaian perbandingan. Average precision baseline adalah 0,576402; scikit-learn 0,736424; Spark 0,736851. ROC-AUC masing-masing 0,5; 0,719782; dan 0,720601. Binary F1 pada ketiganya adalah 0,731288. Precision adalah 0,576402 dan recall adalah 1. Confusion matrix ketiganya adalah TN nol, FP 3.407, FN nol dan TP 4.636. Nilai yang ditampilkan dibulatkan enam desimal; presisi penuh tersedia dalam metrics.json.'),
('Threshold dan batas pelaksanaan',['H1: maksimum F1 validation; ties → threshold terbesar','Baseline: 0.5724722032673705','scikit-learn: 0.3074755313277422','Spark: 0.4466021438573422','Dibekukan sebelum test; prediksi positif p ≥ threshold'],'03:40–04:20','Threshold dipilih pada validation dengan F1 maksimum, termasuk kandidat all-negative. Jika nilainya sama, threshold terbesar dipilih. Angka pada slide dibekukan sebelum prediksi test dan aturan positif adalah probabilitas lebih besar atau sama dengan threshold. Evaluasi kedua jalur memakai evaluator scikit-learn. AP mengacu average_precision_score, bukan penggantian otomatis dengan area trapezoidal precision–recall.'),
('Graph penuh: hasil tercatat',['Bipartit order: / category:; tidak berarah / tidak berbobot','62.948 node = 62.897 pesanan + 51 kategori','152.676 edge; 19 connected components','Self-loop = 0; duplicate edge = 0','Degree Category Id 17 = 19.534; 18 = 17.954; 24 = 17.088'],'04:20–05:10','Graph adalah relasi keanggotaan pesanan dan kategori, bukan rute logistik. Edge unik dibentuk dengan Spark DataFrame distinct, sedangkan degree dan connected components dihitung dengan NetworkX pada satu mesin. Graph penuh memiliki 62.948 node, 152.676 edge dan 19 komponen. Degree kategori menghitung pesanan unik, bukan kuantitas barang. Contoh angka degree Category Id 17 adalah 19.534, Id 18 adalah 17.954, dan Id 24 adalah 17.088. Angka ini tidak disertai klaim kausal atau rekomendasi.'),
('Lingkungan, artefak dan batas bukti',['macOS; Python 3.12.15; Spark 3.5.9 local[2]','file:/// lokal, bukan HDFS; bukan eksekusi Colab','status.json: exit_status=0; training_evaluation_graph=executed','B0: benchmark tidak dilakukan; tanpa klaim speedup','Notebook, parameter, prediksi, metrik dan graph tersedia lokal'],'05:10–05:45','Eksekusi final berjalan pada macOS dengan Python 3.12.15 dan Spark 3.5.9 dalam local dua worker thread. Filesystem adalah file URI lokal, bukan HDFS. Smoke test HDFS historis terpisah. Status final mencatat exit status nol. B0 berarti benchmark tidak dilakukan; cap waktu log bukan bukti percepatan. Notebook tereksekusi, parameter, prediksi, metrik dan hasil graph tersedia secara lokal.'),
('Segmen wajib mahasiswa | bukan naskah AI',['Isi sendiri: alasan fitur, algoritma dan strategi pengembangan','Isi sendiri: interpretasi evaluasi dan perbandingan substantif','Isi sendiri: analisis graph serta kesimpulan','Batas retrospektif, label maturation dan as-of tetap dibahas','Tautan berbagi, deklarasi AI dan video diselesaikan penulis'],'05:45–08:45','[SEGMEN MAHASISWA, MAKSIMAL 3 MENIT: tulis sendiri justifikasi metode, interpretasi hasil, perbandingan substantif, analisis graph dan kesimpulan. Jangan membaca teks placeholder ini sebagai analisis. Gunakan angka aktual pada artefak. Selesaikan dalam slot 3 menit, lalu tutup presentasi. Video teknis diproduksi terpisah; draf teknis AI bukan final mahasiswa.]')
]
prs=Presentation();prs.slide_width=PI(13.333);prs.slide_height=PI(7.5)
for i,(title,bullets,slot,narration) in enumerate(slides,1):
 s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string('F4F6FA')
 def box(x,y,w,h,text,size,color,bold=False):
  shape=s.shapes.add_textbox(PI(x),PI(y),PI(w),PI(h));tf=shape.text_frame;tf.word_wrap=True
  for j,line in enumerate(text.split('\n')):
   p=tf.paragraphs[0] if j==0 else tf.add_paragraph();p.text=line;p.font.name='Aptos';p.font.size=PP(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color);p.space_after=PP(14)
  return shape
 box(.65,.45,12,1.0,title,31,'12304A',True)
 box(.8,1.8,11.8,4.8,'\n'.join(bullets),24,'243748')
 box(.7,7.0,12,.3,f'{i:02d} / 09  |  {slot}  |  Sumber: run-actual; angka dibulatkan 6 desimal',11,'56687A')
 s.notes_slide.notes_text_frame.text=f'{slot}\n{narration}'
ppt=OUT/'Assignment-II-DataCo-Hasil-Aktual.pptx';prs.save(ppt)
script='# Naskah presentasi faktual Assignment II\n\nTarget jadwal: 8 menit 45 detik, termasuk slot mahasiswa 3 menit. Durasi nyata wajib diuji melalui latihan dan dibatasi maksimal 10 menit. Naskah ini bukan video. Angka dibaca dari run-actual; tanpa interpretasi AI.\n\n'
for i,(title,bullets,slot,narration) in enumerate(slides,1):script+=f'## Slide {i} — {title} ({slot})\n\n{narration}\n\n'
script+='## Pekerjaan sebelum penyerahan\n\n- Mahasiswa melengkapi Bagian 5.6 dan 6.2, justifikasi serta segmen 3 menit dengan pemikiran sendiri.\n- Latihan dengan timer; jangan melebihi 10 menit.\n- Unggah notebook/kode/slide dan masukkan tautan yang dapat diakses penilai.\n- Tinjau video teknis yang diproduksi terpisah; draf teknis AI bukan final mahasiswa. Lengkapi segmen mahasiswa dan tautan final.\n- Verifikasi seluruh bantuan AI, prompt, persentase dan deklarasi; tanda tangan oleh mahasiswa.\n- Periksa layout, daftar isi dan maksimal 40 halaman melalui Word/renderer.\n'
(OUT/'Naskah-Presentasi-Faktual.md').write_text(script,encoding='utf-8')
# Actual reopen and package integrity counts.
def counts(path):
 with zipfile.ZipFile(path) as z:
  from lxml import etree
  root=etree.fromstring(z.read('word/document.xml'));ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','m':'http://schemas.openxmlformats.org/officeDocument/2006/math','a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
  return {'active_drawings':len(root.xpath('//w:drawing',namespaces=ns)),'math_objects':len(root.xpath('//m:oMath',namespaces=ns)),'field_chars':len(root.xpath('//w:fldChar',namespaces=ns)),'field_instructions':len(root.xpath('//w:instrText',namespaces=ns)),'active_image_references':len(root.xpath('//a:blip',namespaces=ns)),'media_parts':len([n for n in z.namelist() if n.startswith('word/media/')]),'tables':len(root.xpath('//w:tbl',namespaces=ns))}
a=counts(SRC);b=counts(report)
for k in ['active_drawings','math_objects','field_chars','field_instructions','active_image_references','media_parts']:assert a[k]==b[k],(k,a,b)
reopen=Document(report);rp=Presentation(ppt);assert len(rp.slides)==9;assert all(p.paragraph_format.page_break_before for p in reopen.paragraphs if p.text.startswith('BAB '))
assert before==sha(SRC)
result={'report':str(report),'slides':str(ppt),'script':str(OUT/'Naskah-Presentasi-Faktual.md'),'generator':str(Path(__file__).resolve()),'source_sha256_unchanged':before,'source_counts':a,'output_counts':b,'slides_count':len(rp.slides),'reopened_docx':True,'reopened_pptx':True,'feature_groups':5,'feature_dimensions':48,'schedule_seconds':525,'factual_script_words':sum(len(x[3].split()) for x in slides[:-1]),'render_verified':False,'page_count_verified':False,'toc_refresh_verified':False,'student_analysis_pending':['5.6','6.2','justification and interpretation presentation segment'],'video_available':False,'public_sharing_links_pending':True,'ai_percentage_and_signature':'student only'}
(OUT/'integrity.json').write_text(json.dumps(result,indent=2,ensure_ascii=False));print(json.dumps(result,indent=2,ensure_ascii=False))
