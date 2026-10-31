from pathlib import Path
import json,zipfile,hashlib
from lxml import etree
from docx import Document
from pptx import Presentation
import fitz
P=Path(__file__).parent; R=P.parent.parent; source=R/'laporan-draft/Assignment II - DataCo - Draf Laporan Lengkap.docx'; report=P/'Assignment-II-DataCo-Laporan-Hasil-Aktual.docx'; pdf=P/'Assignment-II-DataCo-Laporan-Hasil-Aktual.pdf'; ppt=P/'Assignment-II-DataCo-Hasil-Aktual.pptx'
def counts(f):
 with zipfile.ZipFile(f) as z:
  x=etree.fromstring(z.read('word/document.xml'));n={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','m':'http://schemas.openxmlformats.org/officeDocument/2006/math','a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
  return {k:len(x.xpath(v,namespaces=n)) for k,v in {'active_drawings':'//w:drawing','math_objects':'//m:oMath','field_chars':'//w:fldChar','field_instructions':'//w:instrText','active_image_references':'//a:blip','tables':'//w:tbl'}.items()} | {'media_parts':len([n for n in z.namelist() if n.startswith('word/media/')])}
a=counts(source);b=counts(report)
for k in ['active_drawings','math_objects','active_image_references','media_parts']:assert a[k]==b[k]
d=Document(report);slides=Presentation(ppt);p=fitz.open(pdf); text='\n'.join(x.text for x in d.paragraphs);pdftext='\n'.join(x.get_text() for x in p)
assert len(p)<=40 and len(slides.slides)==9
assert 'No table of contents entries found.' not in text
assert 'pelatihan Logistic Regression pada' in text
assert 'Tabel 6.3 Komponen graph penuh aktual' in text
assert all(x.paragraph_format.page_break_before for x in d.paragraphs if x.text.startswith('BAB ') and x.style.name=='Heading 1')
for s in slides.slides:
 for sh in s.shapes:
  assert sh.left>=0 and sh.top>=0 and sh.left+sh.width<=slides.slide_width and sh.top+sh.height<=slides.slide_height
# Export a compact page-contact sheet for visual inspection.
thumbs=[]
from PIL import Image,ImageOps,ImageDraw
for i,page in enumerate(p):
 pix=page.get_pixmap(matrix=fitz.Matrix(.35,.35));img=Image.frombytes('RGB',[pix.width,pix.height],pix.samples);img.thumbnail((220,310));canvas=Image.new('RGB',(240,340),'#d9dee5');canvas.paste(img,((240-img.width)//2,15));ImageDraw.Draw(canvas).text((10,320),str(i+1),fill='black');thumbs.append(canvas)
sheet=Image.new('RGB',(240*5,340*6),'white')
for i,im in enumerate(thumbs):sheet.paste(im,((i%5)*240,(i//5)*340))
sheet.save(P/'report-contact-sheet.png')
(P/'report-rendered-text.txt').write_text(pdftext)
(P/'slides-outline.json').write_text(json.dumps([{'slide':i+1,'text':[sh.text for sh in s.shapes if sh.has_text_frame],'notes':s.notes_slide.notes_text_frame.text} for i,s in enumerate(slides.slides)],indent=2,ensure_ascii=False))
result=json.loads((P/'integrity.json').read_text());result.update(output_counts=b,source_counts=a,page_count_verified=True,page_count=len(p),within_40_pages=True,toc_refresh_verified=True,tables_figures_lists_populated=True,pdf_tables_figures_lists_verified=('No table of contents entries found.' not in pdftext),pdf_list_render_blocker='Word PDF retains empty table/figure lists although saved DOCX fields are populated; final PDF not submission-ready',word_fields_updated=73,word_pdf_exported=True,render_verified='Word PDF generated and parsed; visual inspection separately',pptx_structural_bounds_pass=True,slide_visual_render_verified=False,video_available='handled by separate agent; not checked here',report_pdf=str(pdf),contact_sheet=str(P/'report-contact-sheet.png'))
assert hashlib.sha256(source.read_bytes()).hexdigest()==result['source_sha256_unchanged']
(P/'integrity.json').write_text(json.dumps(result,indent=2,ensure_ascii=False));print(json.dumps(result,indent=2,ensure_ascii=False))
