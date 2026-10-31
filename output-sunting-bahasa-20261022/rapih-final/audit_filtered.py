from docx import Document
from pathlib import Path
import json,re
p=Path("/Users/rafaelyudianto/Documents/COMP8035041-Assignment-II/output-sunting-bahasa-20261022/rapih-final/Assignment-II-DataCo-Laporan-Final-Rapi-20261022.docx")
d=Document(p)
items=[]
def scan_para(q,loc):
    text=q.text
    stars=[m.start() for m in re.finditer(r'\*',text)]
    nonempty=[r for r in q.runs if r.text.strip()]
    italic_chars=sum(len(r.text) for r in nonempty if r.italic is True)
    total_chars=sum(len(r.text) for r in nonempty)
    if stars or italic_chars:
        items.append({'loc':loc,'style':q.style.name,'text':text,'stars':len(stars),'italic_chars':italic_chars,'total_chars':total_chars,'all_italic':bool(nonempty) and all(r.italic is True for r in nonempty),'runs':[{'t':r.text,'i':r.italic,'b':r.bold} for r in q.runs if r.text]})
for i,q in enumerate(d.paragraphs): scan_para(q,f'body.p{i}')
for ti,t in enumerate(d.tables):
  for ri,row in enumerate(t.rows):
    for ci,c in enumerate(row.cells):
      for pi,q in enumerate(c.paragraphs): scan_para(q,f'table{ti}.r{ri}.c{ci}.p{pi}')
for si,s in enumerate(d.sections):
  for kind,container in [('header',s.header),('footer',s.footer)]:
    for pi,q in enumerate(container.paragraphs): scan_para(q,f'{kind}{si}.p{pi}')
print(json.dumps([x for x in items if x['stars'] or x['all_italic']],ensure_ascii=False,indent=2))
