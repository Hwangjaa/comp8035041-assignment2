from docx import Document
from pathlib import Path
import re, hashlib, json
p=Path('Assignment-II-DataCo-Laporan-Final-Rapi-20261022.docx'); d=Document(p)
changed=[]
def iter_paras(doc):
    for i,q in enumerate(doc.paragraphs): yield q,f'body.p{i}'
    for ti,t in enumerate(doc.tables):
        for ri,row in enumerate(t.rows):
            for ci,c in enumerate(row.cells):
                for pi,q in enumerate(c.paragraphs): yield q,f'table{ti}.r{ri}.c{ci}.p{pi}'
    for si,s in enumerate(doc.sections):
        for k,container in [('header',s.header),('footer',s.footer)]:
            for pi,q in enumerate(container.paragraphs): yield q,f'{k}{si}.p{pi}'
def fmt(src,dst):
    dst.bold=src.bold; dst.underline=src.underline
    if src.font.name: dst.font.name=src.font.name
    if src.font.size: dst.font.size=src.font.size
for q,loc in list(iter_paras(d)):
    text=q.text
    if '*' not in text or not re.search(r'\*[^*\n]{1,200}\*',text): continue
    src=q.runs[0] if q.runs else None; matches=list(re.finditer(r'\*([^*\n]{1,200})\*',text))
    for r in list(q.runs): r._element.getparent().remove(r._element)
    pos=0
    for m in matches:
        if m.start()>pos:
            r=q.add_run(text[pos:m.start()]);
            if src: fmt(src,r)
            r.italic=False
        r=q.add_run(m.group(1));
        if src: fmt(src,r)
        r.italic=True; pos=m.end()
    if pos<len(text):
        r=q.add_run(text[pos:]);
        if src: fmt(src,r)
        r.italic=False
    changed.append({'loc':loc,'matches':len(matches)})
# Remove accidental all-paragraph italics from the Indonesian body paragraph; retain only English term.
for q,loc in iter_paras(d):
    if loc=='body.p158':
        for r in q.runs:
            if r.text and r.text.strip()!='Preprocessing': r.italic=False
        changed.append({'loc':loc,'action':'reset_nonterm_italic'})
d.save(p)
text='\n'.join(q.text for q,_ in iter_paras(d))
report={'output':str(p.resolve()),'changes':changed,'remaining_literal_asterisks':text.count('*'),'sha256':hashlib.sha256(p.read_bytes()).hexdigest()}
Path('final-audit-report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(report,ensure_ascii=False,indent=2))
