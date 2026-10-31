import json, re, sys, hashlib
from copy import deepcopy
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn

R = Path(__file__).parent
ORIG = R.parent / "Assignment-II-DataCo-Laporan-Final-Rapi-20261022.docx"
TARGET = R / "Assignment-II-DataCo-Revisi-Humanizer.docx"
LOG = R / "rewrite-log.json"

snapshot = {r["i"]: r["text"] for r in json.loads((R / "paragraphs.json").read_text())}

if not TARGET.exists():
    TARGET.write_bytes(ORIG.read_bytes())
doc = Document(TARGET)


def plain(s):
    return s.replace("<i>", "").replace("</i>", "")


def nums(s):
    return sorted(re.findall(r"\d+(?:[.,]\d+)*", plain(s)))


def refs(s):
    return sorted(re.findall(r"(?:Tabel|Gambar|Persamaan)\s*\(?\d+\.\d+", plain(s)))


def pick_rpr(p):
    for r in p.runs:
        rpr = r._r.rPr
        if rpr is not None and rpr.find(qn("w:i")) is None and rpr.find(qn("w:iCs")) is None:
            copy = deepcopy(rpr)
            for tag in ("w:i", "w:iCs"):
                el = copy.find(qn(tag))
                if el is not None:
                    copy.remove(el)
            return copy
    return None


def add_runs(p, text, base_rpr):
    pos = 0
    parts = []
    for m in re.finditer(r"<i>(.*?)</i>", text):
        if m.start() > pos:
            parts.append((text[pos:m.start()], False))
        parts.append((m.group(1), True))
        pos = m.end()
    if pos < len(text):
        parts.append((text[pos:], False))
    for chunk, italic in parts:
        if not chunk:
            continue
        run = p.add_run(chunk)
        if base_rpr is not None:
            run._r.insert(0, deepcopy(base_rpr))
        run.italic = True if italic else False


report = []
for fname in sys.argv[1:]:
    batch = json.loads((R / fname).read_text())
    for key, spec in batch.items():
        i = int(key)
        p = doc.paragraphs[i]
        old = snapshot[i]
        new = spec["new"]
        assert p.text == old, f"p{i}: teks tidak cocok dengan snapshot"
        assert "w:fldChar" not in p._p.xml and "w:instrText" not in p._p.xml, f"p{i}: ada field"
        assert not p._p.xpath(".//w:hyperlink"), f"p{i}: ada hyperlink"
        assert not p._p.xpath(".//m:oMath"), f"p{i}: ada persamaan"
        assert p.text != plain(new), f"p{i}: teks belum berubah"

        if spec.get("number_check") == "spacing":
            assert re.sub(r"\D", "", plain(old)) == re.sub(r"\D", "", plain(new)), f"p{i}: digit berubah"
        else:
            old_nums, new_nums = nums(old), nums(new)
            dropped = sorted(spec.get("drop_numbers", []))
            missing = list(old_nums)
            for n in new_nums:
                if n in missing:
                    missing.remove(n)
            assert sorted(missing) == dropped, f"p{i}: angka hilang {missing} != {dropped}"
            extra = list(new_nums)
            for n in old_nums:
                if n in extra:
                    extra.remove(n)
            allowed_new = sorted(spec.get("add_numbers", []))
            assert sorted(extra) == allowed_new, f"p{i}: angka baru muncul {extra} != {allowed_new}"

        old_refs, new_refs = refs(old), refs(new)
        lost_refs = list(old_refs)
        for r in new_refs:
            if r in lost_refs:
                lost_refs.remove(r)
        assert sorted(lost_refs) == sorted(spec.get("drop_refs", [])), f"p{i}: rujukan hilang {lost_refs}"
        assert not [r for r in new_refs if r not in old_refs], f"p{i}: rujukan baru muncul"

        assert "\u2013" not in new and "\u2014" not in new, f"p{i}: ada en/em dash"
        assert not re.search(r"\b(di mana|yang mana|dimana)\b", plain(new), re.I), f"p{i}: di mana/yang mana"

        base_rpr = pick_rpr(p)
        for child in list(p._p):
            if child.tag != qn("w:pPr"):
                p._p.remove(child)
        add_runs(p, new, base_rpr)
        report.append({"paragraph": i, "batch": fname, "chars_old": len(old), "chars_new": len(plain(new))})

doc.save(TARGET)
prev = json.loads(LOG.read_text()) if LOG.exists() else []
prev.extend(report)
LOG.write_text(json.dumps(prev, ensure_ascii=False, indent=1))
print(json.dumps({"applied": len(report), "total_log": len(prev), "sha256": hashlib.sha256(TARGET.read_bytes()).hexdigest()}, ensure_ascii=False))
