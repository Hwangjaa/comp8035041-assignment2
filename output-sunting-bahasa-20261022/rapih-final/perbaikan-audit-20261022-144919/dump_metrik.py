import argparse
from pathlib import Path
from docx import Document

R = Path(__file__).parent
BAB = {
    "bab1": range(76, 91),
    "bab2": range(91, 109),
    "bab3": range(109, 122),
    "bab4": range(122, 156),
    "bab5": range(156, 192),
    "bab6": range(192, 210),
}

ap = argparse.ArgumentParser()
ap.add_argument("docx")
ap.add_argument("--tag", default="after")
ap.add_argument("--dump", default="")
a = ap.parse_args()
doc = Document(a.docx if Path(a.docx).is_absolute() else R / a.docx)

for name, rng in BAB.items():
    paras = []
    for i in rng:
        p = doc.paragraphs[i]
        if p.style.name.startswith("Heading") or "Caption" in p.style.name:
            continue
        if p.text.strip():
            paras.append(p.text.strip())
    (R / f"metrik-{name}-{a.tag}.md").write_text("\n\n".join(paras), encoding="utf-8")

if a.dump:
    idx = [int(x) for x in a.dump.split(",")]
    for i in idx:
        print(f"[{i}] {doc.paragraphs[i].text}\n")
print("ok")
