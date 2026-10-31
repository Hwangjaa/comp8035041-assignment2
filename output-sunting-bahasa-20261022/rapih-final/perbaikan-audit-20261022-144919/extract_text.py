import json
from pathlib import Path
from docx import Document

R = Path(__file__).parent
ROOT = R.parent
SRC = ROOT / "Assignment-II-DataCo-Laporan-Final-Rapi-20261022.docx"
doc = Document(SRC)

rows = []
for i, p in enumerate(doc.paragraphs):
    rows.append({
        "i": i,
        "style": p.style.name,
        "text": p.text,
        "chars": len(p.text),
        "has_math": bool(p._p.xpath(".//m:oMath")),
        "has_drawing": bool(p._p.xpath(".//w:drawing")),
        "italic_runs": sum(1 for r in p.runs if r.italic),
        "runs": len(p.runs),
    })

(R / "paragraphs.json").write_text(json.dumps(rows, ensure_ascii=False, indent=1), encoding="utf-8")

lines = []
for r in rows:
    tag = "H" if r["style"].startswith("Heading") else ("C" if r["style"] in ("Caption", "Image Caption") else ".")
    if r["text"].strip():
        lines.append(f'[{r["i"]:>3}] {tag} {r["style"][:12]:<12} {r["text"]}')
(R / "paragraphs.txt").write_text("\n".join(lines), encoding="utf-8")
print("paragraf", len(rows), "tabel", len(doc.tables))
print("heading1:")
for r in rows:
    if r["style"] == "Heading 1":
        print(" ", r["i"], r["text"])
