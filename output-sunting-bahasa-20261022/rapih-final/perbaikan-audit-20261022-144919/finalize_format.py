from pathlib import Path
from docx import Document
from lxml import etree
import json,hashlib,re,zipfile
R=Path(__file__).parent
src=R/'Assignment-II-DataCo-Revisi-Audit.docx'
out=R/'Assignment-II-DataCo-Revisi-Format.docx'
d=Document(src)
ps=list(d.paragraphs)+[p for t in d.tables for row in t.rows for c in row.cells for p in c.paragraphs]
changes=[]
for pi,p in enumerate(ps):
    offset=0
    for ri,r in enumerate(p.runs):
        end=offset+len(r.text)
        partial=(offset>0 and p.text[offset-1:offset].isalpha() and r.text[:1].isalpha()) or (end<len(p.text) and p.text[end:end+1].isalpha() and r.text[-1:].isalpha())
        if r.italic and partial and r.text in ('fit','run','log'):
            changes.append({'paragraph':pi,'run':ri,'text':r.text,'context':p.text})
            r.italic=False
        offset=end
assert len(changes)==6,changes
d.save(out)
a=Document(src); b=Document(out)
assert [p.text for p in a.paragraphs]==[p.text for p in b.paragraphs]
assert [[[c.text for c in r.cells] for r in t.rows] for t in a.tables]==[[[c.text for c in r.cells] for r in t.rows] for t in b.tables]
ns={'m':'http://schemas.openxmlformats.org/officeDocument/2006/math','w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
with zipfile.ZipFile(src) as za,zipfile.ZipFile(out) as zb:
    for part in za.namelist():
        if part.startswith('word/media/') or re.match(r'word/(header|footer)',part):assert za.read(part)==zb.read(part),part
    xa=etree.fromstring(za.read('word/document.xml')); xb=etree.fromstring(zb.read('word/document.xml'))
    for path in ('.//m:oMath','.//w:drawing','.//w:fldChar','.//w:instrText'):
        assert [etree.tostring(x) for x in xa.xpath(path,namespaces=ns)]==[etree.tostring(x) for x in xb.xpath(path,namespaces=ns)],path
original=R.parent/'Assignment-II-DataCo-Laporan-Final-Rapi-20261022.docx'
assert hashlib.sha256(original.read_bytes()).hexdigest()=='0d6f9095ba5489cf575afaab8a18bae9ab77181452db142ece8136adc68bc8cb'
report={'output':str(out),'italic_partial_fixed':changes,'text_unchanged_from_editorial':True,'tables':len(b.tables),'equations':len(xb.xpath('.//m:oMath',namespaces=ns)),'drawings':len(xb.xpath('.//w:drawing',namespaces=ns)),'media_equations_headers_footers_fields_unchanged':True,'original_sha256_verified':True,'unresolved':['Rujukan Tabel 6.5 dan 6.7 tidak ada; klaim lateness terkait perlu keputusan substantif, bukan pengalihan nomor tabel.','Klaim repositori, provenance keluaran graf, dan environment perlu remediasi berbasis evidence_remediation.json.','Rekomendasi sitasi di citation_remediation.json belum diterapkan pada varian format ini.','Daftar isi/tabel/gambar merupakan daftar konvensional tanpa nomor halaman; belum dapat disebut daftar isi berhalaman terverifikasi.']}
(R/'format_final_qa.json').write_text(json.dumps(report,indent=2,ensure_ascii=False))
print(json.dumps(report,ensure_ascii=False,indent=2))
