import json, re
from pathlib import Path
from docx import Document

R = Path(__file__).parent
orig = Document(R.parent / "Assignment-II-DataCo-Laporan-Final-Rapi-20261022.docx")
new = Document(R / "Assignment-II-DataCo-Revisi-Humanizer.docx")

# 1. Perbaiki spasi yang rusak pada daftar pustaka (kata tergabung).
fixes = {
    "Estimating impacts oflogistics processes": "Estimating impacts of logistics processes",
    "ACM Transactions on KnowledgeDiscovery from Data": "ACM Transactions on Knowledge Discovery from Data",
    "Journal ofMachine LearningResearch": "Journal of Machine Learning Research",
    "in-memoryclustercomputing": "in-memory cluster computing",
    "Random forests.Machine Learning": "Random forests. Machine Learning",
}
fixed = []
for p in new.paragraphs[210:226]:
    for a, b in fixes.items():
        if a in p.text:
            for r in p.runs:
                if a in r.text:
                    r.text = r.text.replace(a, b)
                    fixed.append((p.text[:40], a))

new.save(R / "Assignment-II-DataCo-Revisi-Humanizer.docx")

# 2. Bandingkan token angka dan rujukan seluruh badan laporan.
def body(doc):
    out = []
    for i, p in enumerate(doc.paragraphs):
        if 75 <= i < 210 and p.text.strip():
            out.append((i, p.text))
    return out

o = dict(body(orig)); n = dict(body(new))
on = sorted(re.findall(r"\d+(?:[.,]\d+)*", " ".join(o.values())))
nn = sorted(re.findall(r"\d+(?:[.,]\d+)*", " ".join(n.values())))

from collections import Counter
co, cn = Counter(on), Counter(nn)
missing = co - cn
added = cn - co

# 3. Cek kata tergabung pada paragraf yang ditulis ulang.
glued = []
for i, t in n.items():
    for w in ("fitur", "runtime", "baseline", "test", "train"):
        pass
    if re.search(r"[a-z]{4}(fitur|runtime|baseline)", t):
        glued.append((i, t[:60]))

italics = []
for i, p in enumerate(new.paragraphs):
    if not p.runs:
        continue
    total = sum(len(r.text) for r in p.runs) or 1
    it = sum(len(r.text) for r in p.runs if r.italic)
    if it / total > 0.5:
        italics.append((i, round(it / total, 2)))

print(json.dumps({
    "angka_hilang": dict(missing),
    "angka_baru": dict(added),
    "kata_tergabung": glued,
    "paragraf_italic_mayoritas": italics,
    "tabel": len(new.tables),
    "persamaan": len(new._element.xpath(".//m:oMath")),
    "gambar": len(new._element.xpath(".//w:drawing")),
    "daftar_pustaka_diperbaiki": len(fixed),
}, ensure_ascii=False, indent=1))
