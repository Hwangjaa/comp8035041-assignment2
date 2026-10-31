import json, re
from pathlib import Path
from docx import Document

R = Path(__file__).parent
P = R / "Assignment-II-DataCo-Revisi-Humanizer.docx"
doc = Document(P)

fixes = {
    "oflogistics": "of logistics",
    "KnowledgeDiscovery": "Knowledge Discovery",
    "ofMachine": "of Machine",
    "LearningResearch": "Learning Research",
    "in-memoryclustercomputing": "in-memory cluster computing",
    "forests.Machine": "forests. Machine",
}
changed = []
for p in doc.paragraphs[210:226]:
    t = p.text
    nt = t
    for a, b in fixes.items():
        nt = nt.replace(a, b)
    if nt != t:
        rpr = p.runs[0]._r.rPr if p.runs and p.runs[0]._r.rPr is not None else None
        from copy import deepcopy
        for child in list(p._p):
            if child.tag.endswith("}pPr"):
                continue
            p._p.remove(child)
        run = p.add_run(nt)
        if rpr is not None:
            run._r.insert(0, deepcopy(rpr))
        changed.append(p.text[:50])
doc.save(P)

bad_italic = []
for i, p in enumerate(doc.paragraphs):
    if not p.runs or not p.text.strip():
        continue
    total = sum(len(r.text) for r in p.runs) or 1
    it = sum(len(r.text) for r in p.runs if r.italic)
    if it / total > 0.9 and not p.style.name.startswith("Heading"):
        bad_italic.append((i, p.text[:50]))

suspicious = []
for i, p in enumerate(doc.paragraphs):
    for m in re.finditer(r"[a-z]{3,}(fitur|runtime|baseline|degree|node|edge|test|training|validation)\b", p.text):
        suspicious.append((i, m.group(0)))

print(json.dumps({
    "pustaka_diperbaiki": len(changed),
    "masih_italic_penuh": bad_italic,
    "kata_tergabung": suspicious[:20],
    "p191": doc.paragraphs[191].text,
}, ensure_ascii=False, indent=1))
