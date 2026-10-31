import json, re, sys
from copy import deepcopy
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn

R = Path(__file__).parent
TARGET = R / "Assignment-II-DataCo-Revisi-Humanizer.docx"
doc = Document(TARGET)


def plain(s):
    return s.replace("<i>", "").replace("</i>", "")


def pick_rpr(p):
    for r in p.runs:
        rpr = r._r.rPr
        if rpr is not None and rpr.find(qn("w:i")) is None and rpr.find(qn("w:iCs")) is None:
            copy = deepcopy(rpr)
            for tag in ("w:i", "w:iCs"):
                el = copy.find(qn(tag))
                if el is not None:
                    copy.remove(el)
            return copy
    return None


def add_runs(p, text, base_rpr):
    pos = 0
    parts = []
    for m in re.finditer(r"<i>(.*?)</i>", text):
        if m.start() > pos:
            parts.append((text[pos:m.start()], False))
        parts.append((m.group(1), True))
        pos = m.end()
    if pos < len(text):
        parts.append((text[pos:], False))
    for chunk, italic in parts:
        if not chunk:
            continue
        run = p.add_run(chunk)
        if base_rpr is not None:
            run._r.insert(0, deepcopy(base_rpr))
        run.italic = True if italic else False


edits = json.loads((R / sys.argv[1]).read_text())
done = []
for key, spec in edits.items():
    i = int(key)
    p = doc.paragraphs[i]
    cur = p.text
    assert cur.startswith(spec["old_start"]), f"p{i}: awalan tidak cocok: {cur[:60]!r}"
    new = spec["new"]
    assert re.sub(r"\D", "", cur) == re.sub(r"\D", "", plain(new)), f"p{i}: digit berubah"
    assert re.findall(r"(?:Tabel|Gambar|Persamaan)\s*\(?\d+\.\d+", cur) == re.findall(
        r"(?:Tabel|Gambar|Persamaan)\s*\(?\d+\.\d+", plain(new)
    ), f"p{i}: rujukan berubah"
    assert "\u2013" not in new and "\u2014" not in new, f"p{i}: ada en/em dash"
    assert cur != plain(new), f"p{i}: tidak berubah"

    base_rpr = None
    for r in p.runs:
        if r._r.rPr is not None:
            base_rpr = r._r.rPr
            break
    for child in list(p._p):
        if child.tag != qn("w:pPr"):
            p._p.remove(child)
    add_runs(p, new, base_rpr)
    done.append(i)

doc.save(TARGET)
print(json.dumps({"refined": done}, ensure_ascii=False))
