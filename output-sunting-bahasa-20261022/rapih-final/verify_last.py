from docx import Document
from pathlib import Path
import json,hashlib
p=Path('Assignment-II-DataCo-Laporan-Final-Rapi-20261022.docx'); d=Document(p)
allp=list(d.paragraphs)+[q for t in d.tables for row in t.rows for c in row.cells for q in c.paragraphs]
print('literal_asterisks',sum(q.text.count('*') for q in allp))
q=d.paragraphs[158]
print('p158_text',q.text[:120])
print('p158_runs',[(r.text,r.italic) for r in q.runs])
print('p158_all_italic',all((r.italic is True) for r in q.runs if r.text.strip()))
print('sha256',hashlib.sha256(p.read_bytes()).hexdigest())
