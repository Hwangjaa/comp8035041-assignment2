
from pathlib import Path
import json, re, shutil, hashlib
from copy import deepcopy
from docx import Document
from docx.text.run import Run
from docx.oxml.ns import qn

ROOT=Path('/Users/rafaelyudianto/Documents/COMP8035041-Assignment-II')
W=ROOT/'output-sunting-bahasa-20261022/work'
SRC=ROOT/'Assignment-II-DataCo-Laporan-Final-High-Quality.docx'
OUT=ROOT/'output-sunting-bahasa-20261022/Assignment-II-DataCo-Laporan-Final-Sunting-Bahasa-20261022.docx'

data=[]
for fn in ['edit-bab1-2.json','edit-bab3-4.json','edit-bab5-6.json']:
    data.append(json.loads((W/fn).read_text()))
aud=json.loads((W/'audit-italic.json').read_text())

def all_paras(doc):
    for p in doc.paragraphs: yield p, 'body'
    for ti,t in enumerate(doc.tables):
        for ri,row in enumerate(t.rows):
            for ci,c in enumerate(row.cells):
                for p in c.paragraphs: yield p, f'table[{ti}].row[{ri}].cell[{ci}]'
    for sec in doc.sections:
        for p in sec.header.paragraphs: yield p,'header'
        for p in sec.footer.paragraphs: yield p,'footer'

doc=Document(SRC)
paras=list(all_paras(doc))
# Apply exact paragraph edits.
edit_report=[]
for d in data:
  for e in d.get('edits',[]):
    # DOCX text omits Markdown asterisks used by source.md for italic spans.
    old_plain=e['old_text'].replace('*','').replace('\\\\','')
    def norm_docx(s): return s.replace('\\\\','').replace('\\','')
    target=norm_docx(old_plain)
    matches=[(p,loc) for p,loc in paras if norm_docx(p.text)==target]
    if len(matches)!=1:
      edit_report.append({'status':'FAIL','old_prefix':e['old_text'][:80],'matches':len(matches)})
      continue
    p,loc=matches[0]
    if p.runs:
      p.runs[0].text=e['new_text']
      for r in p.runs[1:]: r.text=''
    else: p.add_run(e['new_text'])
    edit_report.append({'status':'OK','location':loc,'old':e['old_text'],'new':e['new_text']})

# Recollect after text edits.
paras=list(all_paras(doc))
terms=[]
for d in data: terms += d.get('italics',[])
# Add audit proposals only when not clearly a code/name/feature identifier and not already covered.
skip={'Market','market','Customer Segment','Operations / Logistics','CS Lead','Ingestion','Storage','Processing','Analytics/consumption','Dataset','dataset','Hash','hash','Baseline','baseline','Node','node','Edge','edge','Degree','degree','Refit','training','validation','test','threshold','Threshold','seed','iterations','run','Random Forest','Gradient Boosting','Logistic Regression','precision','recall','binary F1','Confusion matrix','Average precision','Precision','Recall','Binary F1','Graph','Insight','BIG DATA','MACHINE LEARNING','GRAPH ANALYTICS','Security','Compliance','Data Ingestion','Processing Framework','Feature Engineering'}
for x in aud.get('missing_italic',[]):
    t=x['text']
    if t not in skip: terms.append(t)
# Explicit exclusions: identifiers/parameter names and wrongly italicized spans.
wrong=set(x['text'] for x in aud.get('wrongly_italic',[])) | {'zero_division','authentication disabled','seed','iterations'}
terms=sorted({t for t in terms if t and t not in wrong}, key=len, reverse=True)
# Some audit items are exact narrative phrases worth italicizing.
pat=re.compile('|'.join(re.escape(t) for t in terms)) if terms else None
wrongpat=re.compile('|'.join(re.escape(t) for t in sorted(wrong,key=len,reverse=True)))
changed_runs=0
italic_hits=[]

def split_run(run, force_pat):
    txt=run.text
    if not txt: return
    # Identify all candidate spans, with wrong spans taking precedence.
    spans=[]
    for m in force_pat.finditer(txt) if force_pat else []: spans.append((m.start(),m.end(),False,m.group()))
    for m in pat.finditer(txt) if pat else []:
        if not any(a < m.end() and m.start() < b for a,b,_,_ in spans): spans.append((m.start(),m.end(),True,m.group()))
    if not spans: return
    spans.sort(); out=[]; pos=0
    for a,b,it,val in spans:
        if a>pos: out.append((txt[pos:a],run.italic))
        out.append((txt[a:b],it)); pos=b
        italic_hits.append(val)
    if pos<len(txt): out.append((txt[pos:],run.italic))
    parent=run._r.getparent(); idx=parent.index(run._r)
    for text,it in out:
        nr=deepcopy(run._r)
        # remove existing text nodes then set one
        ts=nr.xpath('./w:t')
        if ts: ts[0].text=text
        else:
            from docx.oxml import OxmlElement
            t=OxmlElement('w:t');t.text=text;nr.append(t)
        for t in nr.xpath('./w:t')[1:]: t.getparent().remove(t)
        rp=nr.get_or_add_rPr()
        for tag in ['w:i','w:iCs']:
            el=rp.find(qn(tag))
            if el is not None: rp.remove(el)
        if it:
            from docx.oxml import OxmlElement
            rp.append(OxmlElement('w:i'))
        parent.insert(idx,nr); idx+=1
    parent.remove(run._r)
    return True

for p,loc in paras:
    for r in list(p.runs):
        if split_run(r,wrongpat): changed_runs+=1
# Save.
OUT.parent.mkdir(parents=True,exist_ok=True)
doc.save(OUT)
# Verify edits and counts.
outdoc=Document(OUT); outparas=list(all_paras(outdoc))
remaining=[e['old'] for e in edit_report if e['status']=='OK' and any(p.text==e['old'] for p,_ in outparas)]
# check all requested new texts are present at least once
new_missing=[e['new'] for e in edit_report if e['status']=='OK' and not any(p.text==e['new'] for p,_ in outparas)]
rep={'source':str(SRC),'output':str(OUT),'source_sha256':hashlib.sha256(SRC.read_bytes()).hexdigest(),'output_sha256':hashlib.sha256(OUT.read_bytes()).hexdigest(),'edit_total':len(edit_report),'edit_failures':[x for x in edit_report if x['status']=='FAIL'],'old_text_remaining':len(remaining),'new_text_missing':len(new_missing),'italic_terms':len(terms),'italic_split_runs':changed_runs,'italic_hits':len(italic_hits),'wrong_terms_removed':sorted(wrong)}
(W/'apply-report.json').write_text(json.dumps(rep,ensure_ascii=False,indent=2))
print(json.dumps(rep,ensure_ascii=False,indent=2))
