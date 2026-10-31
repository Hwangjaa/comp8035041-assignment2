from pathlib import Path
from copy import deepcopy
import hashlib,json,re
from docx import Document
from docx.text.paragraph import Paragraph
from docx.text.run import Run
from docx.oxml.ns import qn
R=Path(__file__).resolve().parent.parent
Q=R/'qa-high-quality'
S=R/'eksekusi-final-lokal/deliverables/Assignment-II-DataCo-Laporan-Final-Revisi.docx'
O=R/'Assignment-II-DataCo-Laporan-Final-High-Quality.docx'
d=Document(S)
sourcehash=hashlib.sha256(S.read_bytes()).hexdigest()
changes=[]
def replace_start(start,new):
 p=next(p for p in d.paragraphs if p.text.startswith(start))
 assert not p._p.xpath('.//w:drawing|.//m:oMath|.//w:fldChar|.//w:instrText')
 old=p.text
 first=next((r for r in p.runs if r.text),None)
 props=deepcopy(first._r.rPr) if first is not None and first._r.rPr is not None else None
 p.clear();r=p.add_run(new)
 if props is not None:r._r.insert(0,props)
 r.italic=False
 changes.append({'before':old,'after':new})
replace_start('Rancangan awal pada Assignment I memuat', 'Rancangan awal pada Assignment I memuat agregasi skor item menggunakan nilai maksimum, sebagaimana ditunjukkan pada Persamaan (4.3). Pada persamaan tersebut, R_o adalah skor pesanan o, I_o adalah himpunan item pada pesanan o, dan r_i adalah skor item i. Implementasi Assignment II langsung menghasilkan probabilitas tingkat pesanan sehingga agregasi maksimum skor item tidak digunakan dalam eksperimen.')
replace_start('Persamaan (4.4) menggambarkan', 'Persamaan (4.4) menggambarkan jendela integrasi clickstream selama 24 jam sebelum pesanan pada rancangan Assignment I. Simbol t_order menyatakan waktu pesanan dan t_event menyatakan waktu kejadian clickstream; batas awal disertakan dan waktu pesanan dikecualikan. Integrasi tersebut tidak dijalankan dan tidak menghasilkan fitur pada Assignment II.')
replace_start('Pelatihan dan evaluasi akhir', 'Pelatihan dan evaluasi akhir menghasilkan model, probabilitas, prediksi, serta statistik graf yang tersimpan sebagai artefak eksekusi. Spesifikasi fitur, parameter model, lingkungan percobaan, dan ambang klasifikasi tercatat lengkap sehingga setiap angka hasil dapat ditelusuri. Logistic Regression scikit-learn mencatat 415 iterasi dan Spark MLlib 71 iterasi. Jumlah iterasi bukan ukuran kecepatan karena konfigurasi optimisasi serta regularisasinya berbeda; benchmark waktu dan skalabilitas multi-node tidak dilakukan.')
# Span-level typography: no edits to field codes, OMML, images or bibliography/forms.
terms=['machine learning','graph analytics','market drivers','business drivers','market','batch','streaming','clickstream','filesystem','snapshot','feed','real-time','cluster','thread','pipeline','benchmark','throughput','multi-node','single-node','high availability','authentication disabled','driver','loopback','unstructured','timestamp','encoding','one-hot encoding','multi-hot encoding','multi-hot','training','validation','test','preprocessing','feature engineering','baseline','seed','resampling','tuning','vocabulary','fit','scaling','refit','Average Precision','average precision','precision','recall','binary F1','weighted F1','confusion matrix','true negative','false positive','false negative','true positive','threshold','zero_division','ranking','node','edge','degree','connected components','connected component','self-loop','spring layout','iterations','insight','data ingestion','processing framework','security','compliance','data privacy','data quality']
pattern=re.compile(r'(?<![\w])(?:'+'|'.join(re.escape(t) for t in sorted(terms,key=len,reverse=True))+r')(?![\w])',re.I)
active=False;formats=0
for p in [Paragraph(e,d) for e in d.element.iter(qn('w:p'))]:
 if p.text=='BAB 1. PENDAHULUAN':active=True
 if p.text=='DAFTAR PUSTAKA':active=False
 if not active:continue
 for r in list(p.runs):
  if not r.text or any(x.tag not in (qn('w:rPr'),qn('w:t')) for x in r._r):continue
  text=r.text;spans=[];pos=0
  for m in pattern.finditer(text):
   if m.start()>pos:spans.append((text[pos:m.start()],False))
   spans.append((m.group(),True));pos=m.end()
  if pos<len(text):spans.append((text[pos:],False))
  if not spans:continue
  for val,ital in spans:
   e=deepcopy(r._r);nr=Run(e,p);nr.text=val;nr.italic=ital
   for el in nr._r.xpath('./w:rPr/w:iCs'):el.getparent().remove(el)
   r._r.addprevious(e)
  r._r.getparent().remove(r._r);formats+=1
# Preserve result tables, all headings, diagram and native equations.
a=Document(S)
assert [p.text for p in a.paragraphs if p.style.name.startswith('Heading')]==[p.text for p in d.paragraphs if p.style.name.startswith('Heading')]
assert [[[c.text for c in row.cells] for row in t.rows] for t in a.tables]==[[[c.text for c in row.cells] for row in t.rows] for t in d.tables]
assert len(d.inline_shapes)==2 and len(d.element.xpath('//m:oMath'))==5
d.save(O)
assert hashlib.sha256(S.read_bytes()).hexdigest()==sourcehash
(Q/'edit-audit.json').write_text(json.dumps({'source':str(S),'source_sha256':sourcehash,'output':str(O),'changes':changes,'formatted_text_runs':formats,'tables_unchanged':True,'headings_unchanged':True,'images':2,'native_equations':5},ensure_ascii=False,indent=2))
print(json.dumps({'output':str(O),'changes':len(changes),'formatted_runs':formats}))
