from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from pathlib import Path
import json
base=Path('/Users/rafaelyudianto/Downloads')
d=Document(base/'COMP8035041 - Assignment I - Pembenahan Terverifikasi.docx')
def el(tag,text=None):
 x=OxmlElement('m:'+tag)
 if text is not None:x.text=text
 return x
def run(s):
 x=el('r');pr=el('rPr');sty=el('sty');sty.set(qn('m:val'),'p');pr.append(sty);x.append(pr);x.append(el('t',s));return x
def sub(a,b):
 x=el('sSub');x.append(el('sSubPr'));e=el('e');e.append(run(a));x.append(e);s=el('sub');s.append(run(b));x.append(s);return x
def frac(a,b):
 f=el('f');f.append(el('fPr'))
 for tag,s in [('num',a),('den',b)]:
  x=el(tag);x.append(run(s));f.append(x)
 return f
def eq(after,nodes,num):
 p=OxmlElement('w:p');after._p.addnext(p);para=Paragraph(p,after._parent);para.alignment=1
 op=el('oMathPara');op.append(el('oMathParaPr'));om=el('oMath')
 for node in nodes:om.append(node)
 om.append(run('     ('+num+')'));op.append(om);p.append(op)
 return para
repls={
'Penulis mengusulkan analitik':'Proyek ini menelaah data historis pesanan dan label keterlambatan sebagai dasar analitik sebelum pengiriman. Pemeriksaan yang telah diselesaikan mencakup integritas sumber, struktur data, label, timestamp, serta pembersihan transaksi. Dispatch digunakan sebagai fokus studi, tetapi maknanya tidak disamakan secara otomatis dengan label sumber. Sistem prioritas pesanan belum dibangun dan dampak intervensi operasional belum diukur.',
'Solusi yang dipilih menggunakan':'Pelaksanaan proyek menggunakan Python untuk pemeriksaan dan pembersihan data, serta lingkungan Hadoop/HDFS dan Apache Spark lokal untuk pengujian integrasi infrastruktur. HDFS berjalan single-node; Spark dijalankan dalam mode local[2]. Tahap yang telah dikerjakan dibedakan dari komponen yang belum dilaksanakan.',
'Data historis diperoleh':'Ketiga CSV DataCo versi 5 diunduh dari sumber resmi. Integritasnya diperiksa menggunakan SHA-256 dan seluruh nilai cocok dengan metadata penerbit. Data historis dibaca secara batch menggunakan Python. Ingestion transaksi harian dan streaming clickstream tidak dijalankan pada tahap pelaksanaan yang dilaporkan.',
'Daftar pesanan hanya dapat':'Pengolahan menggunakan snapshot historis, bukan feed transaksi operasional. Tidak ada klaim kesegaran real-time, pembaruan harian produksi, atau respons per detik dalam hasil proyek ini.',
'Clickstream diusulkan':'Pemeriksaan clickstream memperoleh rentang 1 September 2017 pukul 06:00 hingga 31 Januari 2018 pukul 23:58. Struktur, nilai kosong, dan kesamaan baris telah diperiksa. Clickstream belum digabungkan dengan transaksi; manfaat prediktif dan hubungan sebab-akibat tidak dinyatakan sebagai hasil.',
'Data lake merupakan':'Data mentah dan keluaran bersih disimpan pada direktori proyek lokal secara terpisah. HDFS dipasang sebagai backend yang telah dipilih dan diuji menggunakan berkas uji. Dataset lengkap belum dimuat ke HDFS; konversi transaksi ke Parquet dan clickstream ke JSON belum dilakukan. Format aktual berkas proyek adalah CSV.',
'Pemisahan istilah:':'HDFS berfungsi sebagai sistem berkas, sedangkan CSV, Parquet, dan JSON adalah format. Pengujian saat ini hanya membuktikan tulis/baca HDFS dan pembacaan berkas uji melalui Spark. Implementasi cloud belum dilakukan; pelaksanaan yang dilaporkan berlangsung lokal.',
'Skalabilitas rancangan belum':'Pengujian tidak mencakup benchmark, skalabilitas multi-mesin, throughput produksi, atau latensi operasional. Karena itu, hasilnya dibatasi pada keberhasilan pemeriksaan data dan uji integrasi lokal, bukan peningkatan kinerja pada skala Big Data.',
'Integrasi transaksi dan clickstream direncanakan':'Transaksi dan clickstream diperiksa sebagai dua sumber terpisah. Tidak dilakukan join tingkat pesanan atau pelanggan. Aturan integrasi dari desain awal dicatat pada Bagian 4.2 sebagai komponen yang belum dieksekusi, bukan hasil pelaksanaan.',
'Gambar 2.1 Arsitektur':'Gambar 2.1 Arsitektur desain awal dan batas pelaksanaan. Python telah memproses transaksi lokal; HDFS dan Spark telah lolos uji berkas kecil. Jalur Parquet, JSON, streaming, integrasi clickstream, dan model pada diagram belum dieksekusi.',
'Kontrol yang dipilih penulis':'Pada pelaksanaan lokal, layanan HDFS dan koneksi Spark dibatasi ke loopback. Penghapusan empat kolom dilakukan pada CSV bersih, sedangkan sumber mentah dipertahankan untuk reproduksi. Enkripsi, autentikasi produksi, role-based access control, dan pengujian keamanan menyeluruh belum diterapkan. Oleh karena itu, kontrol tersebut tidak diklaim sebagai hasil proyek.',
'Jika terdapat akses tidak sah':'Belum dilakukan simulasi insiden, pengujian pemulihan, atau penegakan retensi otomatis. Prosedur keamanan operasional pada desain awal tetap berada di luar hasil implementasi yang telah diuji.',
'Penulis menggunakan GDPR':'GDPR digunakan sebagai referensi prinsip, bukan penetapan yurisdiksi atau bukti kepatuhan menyeluruh. Empat kolom dihapus dari transaksi bersih, tetapi IP masih terdapat pada sumber clickstream yang disimpan untuk pemeriksaan. Dataset belum dinyatakan anonim dan lapisan analitik clickstream belum dibangun.',
'Pemeriksaan yang direncanakan':'Pemeriksaan yang dijalankan mencakup jumlah baris dan kolom, identitas unik, parsing tanggal, duplikasi penuh, Product Description kosong, kesamaan baris clickstream, serta hubungan label dengan durasi pada cakupan nonbatal. Pemeriksaan kebaruan operasional, perubahan distribusi, rentang seluruh atribut, dan karantina otomatis belum dilakukan. late_days negatif menunjukkan selisih durasi dan tidak otomatis merupakan kesalahan.',
'Kolom pemeriksaan:':'Kolom late_days dihitung sebagai selisih durasi aktual dan terjadwal, dalam satuan hari, sesuai Persamaan (4.1). Penambahannya menghasilkan 50 kolom. Label is_late tetap merupakan alias Late_delivery_risk. Pada seluruh item nonbatal, label konsisten dengan durasi aktual yang melebihi durasi terjadwal. late_days dan durasi aktual tidak digunakan sebagai fitur prediksi sebelum pengiriman.',
'Aturan prioritas yang diusulkan:':'Aturan MAX dari desain awal dituliskan pada Persamaan (4.2). Aturan tersebut belum dieksekusi karena skor risiko item dan model belum tersedia. R_o menyatakan skor prioritas pesanan, r_i skor risiko item, dan I_o himpunan item dalam pesanan o. Nilai MAX bukan probabilitas keterlambatan pesanan yang sudah tervalidasi. Semua item dalam pesanan yang sama memiliki label dan timestamp yang sama; hubungan itu harus dipertimbangkan dalam evaluasi berikutnya.',
'Integrasi clickstream tidak memakai':'Tidak dilakukan pencocokan berdasarkan Order Id, Customer Id, atau IP karena kunci tersebut tidak tersedia bersama. Persamaan (4.3) mendokumentasikan batas jendela 24 jam dari desain awal, bukan integrasi yang sudah dijalankan. t_order adalah waktu pesanan dan t_event waktu kunjungan; batas bawah termasuk dan batas atas tidak termasuk.'
}
for p in list(d.paragraphs):
 for prefix,text in repls.items():
  if p.text.startswith(prefix):p.text=text;break
 if p.text.startswith('Pemeriksaan ulang menghasilkan'):
  p.text='Pemeriksaan memperoleh 172.765 item nonbatal dalam 62.897 pesanan. Sebanyak 98.977 item berlabel 1 dan 73.788 berlabel 0. Persentase item terlambat dihitung pada Persamaan (1.1), bukan persentase pesanan atau pelanggan. Sebanyak 7.754 item Shipping canceled dikeluarkan sesuai aturan cakupan.'
  eq(p,[run('P = '),frac('98.977','172.765'),run(' × 100% ≈ 57,29%')],'1.1')
 if p.text.startswith('Pemeriksaan memperoleh 180.519'):
  p.text='Data sumber memuat 180.519 Order Item Id unik dan 65.752 Order Id unik. Rata-rata item per pesanan dihitung pada Persamaan (4.0). Cakupan nonbatal memuat 62.897 pesanan. Label dan kedua timestamp tidak bervariasi antaritem dalam setiap Order Id.'
  eq(p,[run('μ = '),frac('180.519','65.752'),run(' ≈ 2,75 item/pesanan')],'4.0')
 if p.text.startswith('Kolom late_days dihitung'):eq(p,[sub('d','late'),run(' = '),sub('d','actual'),run(' − '),sub('d','scheduled')],'4.1')
 if p.text.startswith('Aturan MAX dari desain'):
  lim=el('limLow');e=el('e');e.append(run('max'));lim.append(e);l=el('lim');l.append(run('i ∈ I'));l.append(sub('','o'));lim.append(l)
  eq(p,[sub('R','o'),run(' = '),lim,run(' '),sub('r','i')],'4.2')
 if p.text.startswith('Tidak dilakukan pencocokan'):eq(p,[sub('t','order'),run(' − 24 h ≤ '),sub('t','event'),run(' < '),sub('t','order')],'4.3')
# Hindari rumus teks polos dalam tabel; angka jejak perubahan tetap tersedia.
for t in d.tables:
 for row in t.rows:
  for c in row.cells:
   s=c.text
   if s=='49 = 53 − 4':c.text='Empat kolom dihapus'
   elif s=='50 = 49 + 1':c.text='Satu kolom late_days ditambahkan'
   elif s=='172.765 = 180.519 − 7.754':c.text='7.754 item batal dikeluarkan'
   elif s=='Aritmetika perubahan jumlah':c.text='Jumlah keluaran bersih'
   elif s=='Tidak mengonfirmasi isi CSV':c.text='CSV dibaca ulang; 180.519 dan 172.765 baris, 50 kolom'
   elif s=='3.249 baris identik dilaporkan':c.text='3.249 baris identik terverifikasi'
   elif s=='Kondisi yang dilaporkan':c.text='Pelaksanaan terverifikasi'
   elif s=='Tabel 2.1 Implementasi dan rancangan operasional':c.text='Tabel 2.1 Pelaksanaan dan komponen belum dieksekusi'
# Rapikan daftar caption; tidak pakai nomor halaman lama.
for p in d.paragraphs:
 if p.text.startswith('\tTabel ') or p.text.startswith('\tGambar '):
  p.text=p.text.rsplit('\t',1)[0].lstrip('\t')
  if p.text.startswith('Gambar 2.1'):p.text='Gambar 2.1 Arsitektur desain awal dan batas pelaksanaan'
 if p.text=='Tabel 2.1 Implementasi dan rancangan operasional':p.text='Tabel 2.1 Pelaksanaan dan komponen belum dieksekusi'
# Word rendra les expressions OMML en mathématiques natives éditables.
path=base/'COMP8035041 - Assignment I - Narasi Pelaksanaan dan Equation.docx'
d.save(path)
check=Document(path);n=len(check.element.xpath('.//m:oMath'))
assert n==5,n
assert len(check.element.xpath('.//m:f'))==2
assert len(check.tables)==12 and len(check.inline_shapes)==1
print(json.dumps({'file':str(path),'native_equations':n,'fractions':2,'tables':len(check.tables),'images':len(check.inline_shapes),'checks':'PASS'},indent=2))
