from pathlib import Path
from copy import deepcopy
import json,re
from docx import Document
from docx.shared import Inches,Pt
from docx.enum.text import WD_TAB_ALIGNMENT,WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
from PIL import Image,ImageDraw,ImageFont
R=Path('/Users/rafaelyudianto/Documents/COMP8035041-Assignment-II');B=Path('/Users/rafaelyudianto/Downloads')
src=Document(B/'COMP8035041 - Assignment I - Gambar dan Daftar Final.docx')
# White-background diagrams; each box carries a status, not color alone.
font='/System/Library/Fonts/Supplemental/Arial.ttf';ital='/System/Library/Fonts/Supplemental/Arial Italic.ttf'
def diagram(name,title,items,notes):
 w=1700;h=180+len(items)*185+len(notes)*40+80
 im=Image.new('RGB',(w,h),'white');dr=ImageDraw.Draw(im)
 def text(t,y,size=30):
  f=ImageFont.truetype(ital if any(s in t for s in ['Spark','HDFS','CSV','Python','Parquet','JSON','streaming','local[2]']) else font,size)
  bb=dr.textbbox((0,0),t,font=f);dr.text(((w-bb[2])/2,y),t,font=f,fill='#172b4d')
 text(title,30,38)
 colors={'UJI':('#e9f5ed','#26734a'),'BELUM':('#fff5df','#956a12'),'MASA DEPAN':('#eef0f5','#68728a')}
 for i,(status,label,line1,line2) in enumerate(items):
  y=115+i*185;fill,stroke=colors[status]
  if i:
   dr.line((850,y-38,850,y-10),fill='#526074',width=4);dr.polygon([(850,y-3),(840,y-17),(860,y-17)],fill='#526074')
  dr.rounded_rectangle((70,y,1630,y+145),radius=14,fill=fill,outline=stroke,width=3)
  text('['+status+'] '+label,y+12,32);text(line1,y+60,27);text(line2,y+100,26)
 y=130+len(items)*185
 for i,t in enumerate(notes):text(t,y+i*40,24)
 out=R/name;im.save(out);return out
img1=diagram('architecture-revised.png','Arsitektur Konseptual DataCo: Penyimpanan dan Pemrosesan',[
 ('UJI','SUMBER — snapshot historis CSV','Transaksi 180.519 × 53; clickstream 469.977 × 8','Unduhan, SHA-256 dan pemeriksaan struktur telah dilakukan'),
 ('BELUM','INGESTION — pemasukan batch ke HDFS','Pengumpulan CSV dan audit Python lokal sudah selesai','Pemuatan seluruh dataset ke HDFS belum dilakukan'),
 ('BELUM','STORAGE — HDFS dalam ekosistem Apache Hadoop','Lapisan mentah → hasil pembersihan → siap analitik','HDFS single-node diuji; lapisan dataset lengkap belum dimuat'),
 ('BELUM','PROCESSING — Apache Spark','Baca HDFS → transformasi → agregasi → tulis hasil','Spark local[2] membaca tiga baris uji; transformasi penuh belum'),
 ('BELUM','ANALYTICS — eksperimen dan skor risiko','Target tetap label sumber; makna dispatch belum tervalidasi','Model, evaluasi, dan skor prioritas belum tersedia'),
 ('BELUM','CONSUMPTION — penggunaan keluaran','Daftar prioritas untuk Operations/Logistics dan CS Lead','Keluaran usulan; bukan sistem operasional yang telah digunakan')],['CSV: format aktual. Parquet transaksi dan JSON clickstream: format target, belum dikonversi.','Streaming hanya pengembangan masa depan; membutuhkan sumber berkelanjutan dan uji tersendiri.','Warna hijau = bukti tahap sumber/uji; kuning = implementasi dataset lengkap belum dilakukan.'])
img2=diagram('spark-flow-revised.png','Alur Apache Spark dan HDFS: Bukti Uji vs Pipeline Dataset',[
 ('UJI','JALUR UJI — berkas kecil','Tulis/baca berkas di HDFS → Spark local[2] → tiga baris','HDFS_WRITE_READ=PASS; SPARK_HDFS_SMOKE=PASS'),
 ('BELUM','JALUR DATASET — CSV historis → HDFS mentah','Pemasukan batch snapshot transaksi dan clickstream','Tidak ada bukti pemuatan seluruh dataset ke HDFS'),
 ('BELUM','SPARK — pembacaan dan transformasi transaksi','Pemetaan kolom → parsing tanggal → pemeriksaan → cakupan nonbatal','Aturan telah dieksekusi di Python lokal, belum melalui Spark'),
 ('BELUM','SPARK ↔ HDFS — hasil pembersihan dan agregasi','Transaksi bersih ke Parquet; agregasi sesuai tujuan analitik','Konversi, pembacaan ulang hasil, dan agregasi Spark belum diuji'),
 ('BELUM','SPARK — eksperimen machine learning','Data siap analitik → model → evaluasi → skor item/pesanan','Algoritma, fitur dan evaluasi belum final; tidak ada hasil model'),
 ('MASA DEPAN','CLICKSTREAM — jalur terpisah','JSON target; integrasi agregat waktu/produk belum dilakukan','Structured Streaming bukan ingestion otomatis dari CSV historis')],['Apache Hadoop mendukung HDFS; MapReduce bukan mesin pemrosesan proyek.','YARN tidak digunakan: Spark berjalan lokal dengan dua thread, bukan dua node.'])
# Keep source document styles, cover and verified body; discard all broken front matter.
d=Document(B/'COMP8035041 - Assignment I - Gambar dan Daftar Final.docx')
body=d.element.body
start=next(p._p for p in d.paragraphs if p.text=='DAFTAR ISI')
end=next(p._p for p in d.paragraphs if p.text=='BAB 1. PENDAHULUAN')
n=start
while n is not end:
 nxt=n.getnext();body.remove(n);n=nxt
# Clean all chapter pagination redundancies; chapter heads have explicit pageBreakBefore.
for p in list(d.paragraphs):
 if not p.text and p._p.xpath('.//w:br[@w:type="page"]'):p._p.getparent().remove(p._p)
# Paragraph replacements address architecture and evidence limits without new decisions.
updates={
'Pelaksanaan proyek menggunakan':'Arsitektur menggunakan Apache Hadoop sebagai ekosistem pendukung, HDFS sebagai komponen penyimpanan, dan Apache Spark sebagai mesin pemrosesan utama. HDFS bukan pengganti Hadoop, melainkan bagian dari ekosistemnya. MapReduce tidak digunakan untuk transformasi atau eksperimen proyek. Pengumpulan historis, penyimpanan, pemrosesan, analitik, dan penggunaan keluaran dibedakan agar keberhasilan satu tahap tidak dianggap membuktikan seluruh pipeline.',
'Data mentah dan keluaran bersih':'HDFS dipilih untuk memisahkan data mentah, hasil pembersihan, dan data siap analitik dalam satu namespace yang dapat dibaca Spark. Manfaatnya pada proyek ini adalah pembelajaran interaksi storage–processing dan keterlacakan lapisan data, bukan peningkatan kinerja yang sudah terbukti. Struktur berlapis adalah keputusan desain; yang benar-benar tersedia saat ini ialah direktori dan CSV lokal serta berkas uji pada HDFS. Dataset lengkap belum dimuat ke HDFS.',
'HDFS berfungsi':'HDFS menggunakan NameNode untuk metadata dan DataNode untuk blok data. Konfigurasi lokal memakai satu DataNode dan dfs.replication=1. Karena tidak ada salinan blok lintas mesin, pengujian ini tidak membuktikan toleransi kegagalan node, high availability, atau skalabilitas produksi. Snapshot CSV mendukung pola baca batch. Parquet untuk transaksi bersih dan JSON untuk clickstream merupakan format target, bukan keluaran konversi yang telah tersedia.',
'Apache Spark 3.5.9':'Apache Spark 3.5.9 berjalan dengan distribusi Apache Hadoop 3.4.3 yang menyediakan HDFS, menggunakan Java 17.0.20.1. Spark local[2] berarti dua thread eksekusi lokal, bukan dua node. Uji Java melalui spark-submit berhasil membaca tiga baris dari hdfs://127.0.0.1:19000/smoke/input.txt. YARN tidak digunakan karena belum ada kebutuhan pengelolaan resource cluster. Tidak ada klaim pipeline PySpark dataset lengkap sudah berjalan.',
'Pengujian tidak mencakup':'Spark dipilih sebagai framework utama karena dapat menjalankan transformasi tabular, agregasi, dan mendukung eksperimen machine learning dalam ekosistem pemrosesan yang sama. Kemampuan teknologi tersebut adalah dasar desain; keberhasilan pekerjaan proyek tetap memerlukan eksekusi tersendiri. Dataset saat ini dapat diperiksa lokal. Belum tersedia benchmark waktu, penggunaan memori, throughput, atau skalabilitas multi-node, sehingga HDFS dan Spark tidak dinyatakan lebih cepat daripada pengolahan lokal.',
'Transaksi dan clickstream diperiksa':'Gambar 2.1 memisahkan fungsi sumber, ingestion, storage, processing, analytics, dan consumption. Gambar 2.2 memperbesar jalur HDFS–Spark: CSV historis masuk secara batch ke HDFS, Spark membaca dan mentransformasi, lalu menulis keluaran pada lapisan hasil pembersihan dan siap analitik. Alur ini adalah desain yang belum dieksekusi penuh. Jalur yang terverifikasi hanya berkas uji kecil; pembersihan dataset aktual dilakukan melalui Python csv di direktori lokal.',
'Gambar 2.2 memperlihatkan':'Pada jalur dataset, transformasi mengikuti aturan yang sudah dipakai di Python: pemetaan nama kolom, parsing tanggal, penghapusan empat atribut, pembentukan late_days untuk audit, dan pemilihan cakupan nonbatal. Agregasi risiko belum dijalankan karena model belum menghasilkan skor. Jalur streaming memerlukan sumber kejadian masuk berkelanjutan dan pengujian terpisah; CSV historis tidak membuktikan layanan streaming operasional.',
'Pada pelaksanaan lokal':'Konfigurasi HDFS menetapkan alamat loopback untuk NameNode dan DataNode. Log Spark menunjukkan koneksi driver pada 127.0.0.1 dan authentication disabled. Pembatasan jaringan ini mengurangi paparan eksternal tetapi tidak menggantikan autentikasi atau otorisasi. Enkripsi data, kontrol akses produksi, dan audit keamanan menyeluruh belum dibuktikan. Penghapusan empat atribut dari keluaran bersih tidak menjadikan seluruh sumber anonim.',
'Sumber resmi menyatakan':'DataCo versi 5 dirilis dengan CC BY 4.0. Atribusi menggunakan Constante, Silva, dan Pereira (2019), DOI 10.17632/8gx2fvg2k6.5. Sumber resmi: https://data.mendeley.com/datasets/8gx2fvg2k6/5. Lisensi: https://creativecommons.org/licenses/by/4.0/. Perubahan mencakup pemetaan kolom, konversi representasi tanggal, penghapusan empat kolom, late_days, dan cakupan nonbatal. Lisensi tidak membuktikan anonimisasi ataupun kepatuhan hukum menyeluruh.',
'Pengecualian kolom hasil akhir':'Durasi aktual, shipping_date, late_days, Delivery Status, dan status akhir lainnya dilarang sebagai masukan prediksi sebelum pengiriman. Late_delivery_risk/is_late adalah label, bukan fitur. Order Status, profit, serta atribut tersisa perlu ditelusuri waktu ketersediaannya; tidak otomatis aman karena tidak tercantum pada daftar larangan. Matriks fitur belum dibangun, sehingga belum dapat disahkan bebas data leakage. Pembersihan dan label audit tidak sama dengan penyiapan masukan model. Hubungan antaritem dalam Order Id yang sama perlu dijaga dalam evaluasi; strategi split dan metrik belum diputuskan.',
'Data diorganisasi menjadi':'Organisasi aktual memisahkan data/raw berisi tiga CSV, data/clean berisi dua CSV transaksi, serta data/audit-results.json dan data/source-manifest.json. HDFS saat ini menyimpan berkas uji di /smoke/input.txt. Lapisan HDFS untuk dataset mentah, bersih, dan siap analitik adalah rancangan yang belum dibuktikan berisi dataset. CSV audit masih memuat label dan atribut hasil akhir; belum merupakan matriks fitur.',
'Ketepatan waktu proses logistik':'Akturk et al. (2022) melaporkan hubungan durasi pemrosesan dan pengiriman dengan penilaian pelanggan. Temuan literatur mendukung relevansi masalah, bukan pengukuran dampak bisnis DataCo. Proyek ini belum mengukur penurunan biaya, kepuasan pelanggan, atau keberhasilan intervensi. Tujuan analitik tetap penyediaan informasi risiko sebelum pengiriman; skor dan daftar prioritas masih keluaran yang diusulkan.',
'Kelanjutan Assignment II':'Kelanjutan Assignment II memerlukan penentuan data dan metode oleh penulis, eksekusi Python untuk model dengan dan tanpa Spark MLlib atau library relevan, evaluasi, serta graph analytics. Laporan ini berhenti pada data, desain, dan uji infrastruktur yang tersedia; tidak mencantumkan metrik atau kesimpulan eksperimen yang belum dilakukan.'}
for p in d.paragraphs:
 for prefix,text in updates.items():
  if p.text.startswith(prefix):p.text=text;break
# Replace diagrams in place, preserving caption adjacency.
shapes=list(d.inline_shapes)
for shape,path in zip(shapes,[img1,img2]):
 par=shape._inline.getparent().getparent();parent=par.getparent();index=parent.index(par)
 node=OxmlElement('w:p');parent.insert(index,node);pp=Paragraph(node,d._body);pp.alignment=1;pp.paragraph_format.keep_with_next=True
 pp.add_run().add_picture(str(path),width=Inches(6.4));parent.remove(par)
for p in d.paragraphs:
 if p.text.startswith('Gambar 2.1') and p.style.name=='Figure Caption':p.text='Gambar 2.1 Arsitektur konseptual Hadoop–HDFS–Spark dan status implementasi'
 if p.text.startswith('Gambar 2.2') and p.style.name=='Figure Caption':p.text='Gambar 2.2 Alur pemrosesan HDFS–Spark dan batas bukti eksekusi'
# Tables: rebuild rows only where stale design/implementation semantics conflict.
def reset_table(table,headers,rows):
 for row in list(table.rows)[1:]:table._tbl.remove(row._tr)
 for c,t in zip(table.rows[0].cells,headers):c.text=t
 for values in rows:
  for c,t in zip(table.add_row().cells,values):c.text=t
reset_table(d.tables[2],['Fungsi','Bukti pelaksanaan','Belum dilaksanakan'],[
 ['Ingestion','Unduh snapshot CSV; verifikasi hash; baca batch Python lokal','Pemuatan seluruh CSV ke HDFS; batch terjadwal'],['Storage','HDFS single-node; tulis/baca berkas uji','Lapisan dataset lengkap; konversi Parquet/JSON'],['Processing','Pembersihan Python; Spark membaca tiga baris HDFS','Transformasi dan agregasi dataset melalui Spark'],['Analytics/consumption','Belum ada skor atau daftar prioritas','Model, evaluasi, prioritas, dan penggunaan operasional']])
reset_table(d.tables[3],['Komponen','Kondisi aktual','Status desain'],[['Apache Hadoop / HDFS','Hadoop 3.4.3; HDFS single-node; replication=1','Backend lapisan data lake; bukan alternatif Hadoop'],['CSV mentah/bersih','Berkas lokal tersedia dan diperiksa','Dataset lengkap belum dimuat ke HDFS'],['Parquet transaksi','Tidak ada konversi terverifikasi','Format target hasil pembersihan'],['JSON clickstream','Sumber aktual masih CSV','Format target; bukan hasil tersedia']])
reset_table(d.tables[8],['Lapisan','Isi','Lokasi/status'],[['Mentah','Tiga CSV dan manifest SHA-256','data/raw lokal; terverifikasi'],['Bersih','dataco_clean_items.csv; dataco_risk_scope.csv','data/clean lokal; 50 kolom'],['Siap analitik','Fitur dan label terpisah','Belum dibangun'],['HDFS','/smoke/input.txt','Berkas uji; lapisan dataset lengkap belum dimuat']])
reset_table(d.tables[10],['Sumber/keluaran','Unit/model','Operasi','Status'],[['Transaksi','Tabular/item','Pemetaan, tanggal, empat kolom, late_days, nonbatal','Python csv lokal; terverifikasi'],['Clickstream','Tabular/kunjungan','Struktur, timestamp, kosong, identik penuh','Audit selesai; IP tetap sumber; deduplikasi belum'],['Integrasi','Agregat waktu/produk','Jendela 24 jam pra-pesanan','Desain; mapping dan pasangan belum'],['Prioritas','Order Id','MAX skor risiko item','Desain; model/skor belum ada'],['Aliran masa depan','Kejadian berkelanjutan','Structured Streaming','Pengembangan; belum diuji']])
# Repair broken bibliographic entries without inventing metadata; preserve unused sources transparently in audit.
ps=list(d.paragraphs)
for a,b in zip(ps,ps[1:]):
 if a.text.startswith('Kaufman,') and b.text.startswith('Formulation,'):a.text=a.text+' '+b.text;b._p.getparent().remove(b._p)
 if a.text.startswith('Zaharia,') and b.text.startswith('Shenker,'):a.text=a.text+' '+b.text;b._p.getparent().remove(b._p)
# Only cited source entries stay in main references; removed entries retained in original and audit.
keep=('Akturk,','Constante,','European Parliament','Kapoor,','Kaufman,','Zaharia,')
inrefs=False
removed=[]
for p in list(d.paragraphs):
 if p.text=='DAFTAR PUSTAKA':inrefs=True;continue
 if p.text=='BINUS UNIVERSITY ARTIFICIAL INTELLIGENCE USE DECLARATION FORM':inrefs=False
 if inrefs and p.text and not p.text.startswith(keep+('Video presentasi:',)):
  removed.append(p.text);p._p.getparent().remove(p._p)
# Add official docs; versions retrieved differ from installed, noted explicitly.
anchor=next(p for p in d.paragraphs if p.text.startswith('Video presentasi:'))
for text in ['Apache Software Foundation. HDFS Architecture. Dokumentasi current (halaman yang diperiksa berlabel 3.5.0); rujukan konsep, bukan bukti versi lokal. https://hadoop.apache.org/docs/current/hadoop-project-dist/hadoop-hdfs/HdfsDesign.html','Apache Software Foundation. Submitting Applications. Spark 3.5.7 Documentation; rujukan mode eksekusi, bukan bukti versi lokal 3.5.9. https://spark.apache.org/docs/3.5.7/submitting-applications.html']:
 node=OxmlElement('w:p');anchor._p.addprevious(node);Paragraph(node,anchor._parent).add_run(text)
# Add bounded citations and traceability notes in relevant chapters.
for p in d.paragraphs:
 if p.text.startswith('HDFS menggunakan NameNode'):p.add_run(' Konsep mengikuti dokumentasi Apache HDFS; kondisi konfigurasi mengacu pada hdfs-site.xml lokal.')
 if p.text.startswith('Spark dipilih sebagai'):p.add_run(' Rujukan dasar pemrosesan Spark: Zaharia et al. (2012); mode eksekusi dijelaskan dokumentasi Apache Spark.')
# Clean front matter: conventional uninterrupted entries, not headings.
anchor=next(p for p in d.paragraphs if p.text=='BAB 1. PENDAHULUAN')
entries=[p.text for p in d.paragraphs if p.style.name.startswith('Heading') and (p.text.startswith('BAB ') or re.match(r'\d+\.\d+',p.text) or p.text in ['DAFTAR PUSTAKA','BINUS UNIVERSITY ARTIFICIAL INTELLIGENCE USE DECLARATION FORM'])]
fronts=[('DAFTAR ISI',['DAFTAR TABEL','DAFTAR GAMBAR']+entries),('DAFTAR TABEL',[p.text for p in d.paragraphs if p.style.name=='Table Caption']),('DAFTAR GAMBAR',[p.text for p in d.paragraphs if p.style.name=='Figure Caption'])]
for title,items in fronts:
 node=OxmlElement('w:p');anchor._p.addprevious(node);p=Paragraph(node,anchor._parent);p.style='Heading 1';p.text=title;p.paragraph_format.page_break_before=True
 for text in items:
  node=OxmlElement('w:p');anchor._p.addprevious(node);p=Paragraph(node,anchor._parent);p.style='Normal';p.paragraph_format.page_break_before=False;p.paragraph_format.keep_with_next=False;p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1
  p.paragraph_format.tab_stops.add_tab_stop(Inches(6.5),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS);p.add_run(text+'\t0')
# Formatting all new content; equation OMML untouched.
foreign=re.compile(r'Apache Hadoop|Apache Spark|Structured Streaming|Spark MLlib|Processing Framework|Data Ingestion|Market|Business Drivers|Security|Compliance|machine learning|data leakage|data lake|clickstream|dispatch|batch|storage|processing|analytics|consumption|ingestion|local\[2\]|single-node|multi-node|thread|NameNode|DataNode|namespace|replication|snapshot|framework|pipeline|benchmark|throughput|high availability|resource|cluster|MapReduce|YARN|HDFS|Hadoop|Spark|Java|Python|Parquet|JSON|CSV|SHA-256|GDPR|Order Id|Order Item Id|Late_delivery_risk|Delivery Status|Order Status|shipping_date|late_days|is_late|profit|split|MAX|Operations|Logistics|CS Lead|graph analytics|spark-submit|authentication disabled|loopback|IP',re.I)
paras=list(d.paragraphs)
for table in d.tables:
 for i,row in enumerate(table.rows):
  pr=row._tr.get_or_add_trPr()
  if not pr.xpath('./w:cantSplit'):pr.append(OxmlElement('w:cantSplit'))
  for cell in row.cells:
   paras.extend(cell.paragraphs)
   for p in cell.paragraphs:
    p.paragraph_format.line_spacing=1;p.paragraph_format.space_after=Pt(3)
    for run in p.runs:run.font.name='Times New Roman';run.font.size=Pt(9);run.bold=(i==0)
for p in paras:
 for run in list(p.runs):
  if run._r.xpath('.//w:drawing'):continue
  matches=list(foreign.finditer(run.text))
  if not matches:continue
  text=run.text;pos=0
  for m in matches:
   for part,it in [(text[pos:m.start()],False),(m.group(),True)]:
    if not part:continue
    rr=deepcopy(run._r)
    for ch in list(rr):
     if ch.tag!='{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rPr':rr.remove(ch)
    tt=OxmlElement('w:t');tt.set('{http://www.w3.org/XML/1998/namespace}space','preserve');tt.text=part;rr.append(tt);run._r.addprevious(rr)
    if it:
     from docx.text.run import Run
     Run(rr,p).italic=True
   pos=m.end()
  if text[pos:]:
   rr=deepcopy(run._r)
   for ch in list(rr):
    if ch.tag!='{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rPr':rr.remove(ch)
   tt=OxmlElement('w:t');tt.set('{http://www.w3.org/XML/1998/namespace}space','preserve');tt.text=text[pos:];rr.append(tt);run._r.addprevious(rr)
  run._r.getparent().remove(run._r)
for p in d.paragraphs:
 if p.text.startswith('BAB ') or p.text in ['DAFTAR PUSTAKA','BINUS UNIVERSITY ARTIFICIAL INTELLIGENCE USE DECLARATION FORM']:p.paragraph_format.page_break_before=True
 if p.style.name.startswith('Heading'):p.paragraph_format.keep_with_next=True
# Exact required title; do not overwrite source.
out=B/'Assignment I - Big Data Architecture Revised.docx';d.save(out)
# Persist independent findings and evidence map, including excluded unverified references.
a={'source':str(B/'COMP8035041 - Assignment I - Gambar dan Daftar Final.docx'),'output':str(out),'evidence':['data/audit-results.json','data/source-manifest.json','audit_data.py','runtime/conf/hadoop/hdfs-site.xml','runtime/logs/smoke-test.log','runtime/smoke/HdfsSparkSmoke.java'],'closed':['Hadoop/HDFS relationship','storage-processing separation','version attribution','two diagrams in body','duplicate front matter','format actual vs target','single-node replication limitation','leakage status not falsely cleared'],'open':['semantic dispatch evidence','feature availability and evaluation design','full HDFS dataset load','Spark transformations and conversions','clickstream mapping and integration','model and priority evaluation','streaming and benchmarks','AI final contribution and author approval','some retained reference metadata not fully verified'],'excluded_reference_entries':removed,'official_docs_version_caveat':'Concept references current HDFS 3.5.0 / Spark 3.5.7; installed versions proven locally, not inferred from docs.'}
(R/'architecture-final-audit.json').write_text(json.dumps(a,ensure_ascii=False,indent=2))
c=Document(out);assert len(c.inline_shapes)==2 and len(c.element.xpath('.//m:oMath'))==5
print(json.dumps({'output':str(out),'images':2,'equations':5,'tables':len(c.tables),'audit':str(R/'architecture-final-audit.json')},indent=2))
