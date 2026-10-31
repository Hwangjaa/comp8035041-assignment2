from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches,Pt
from docx.enum.text import WD_TAB_ALIGNMENT,WD_TAB_LEADER
OUT=Path(__file__).resolve().parent
p=OUT/'Assignment-II-DataCo-Revisi-QA-v2.docx'
d=Document(p)
# Daftar berisi judul caption + field PAGEREF dan hyperlink internal nyata.
for label,style,heading in [('tab','Table Caption','DAFTAR TABEL'),('fig','Figure Caption','DAFTAR GAMBAR')]:
    captions=[x for x in d.paragraphs if x.style.name==style]
    for i,cap in enumerate(captions,1):
        name=f'qa_{label}_{i}';start=OxmlElement('w:bookmarkStart');start.set(qn('w:id'),str(5000+(100 if label=='fig' else 0)+i));start.set(qn('w:name'),name)
        end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),start.get(qn('w:id')))
        cap._p.insert(0,start);cap._p.append(end)
    h=next(x for x in d.paragraphs if x.text==heading)
    old=h._p.getnext()
    # paragraf TOC kosong lama dengan field yang tidak terindeks
    if old is not None and old.xpath('.//w:instrText'):
        old.getparent().remove(old)
    anchor=h
    for i,cap in enumerate(captions,1):
        para=d.add_paragraph();anchor._p.addnext(para._p);anchor=para
        para.paragraph_format.tab_stops.add_tab_stop(Inches(6.45),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
        para.paragraph_format.space_after=Pt(6)
        link=OxmlElement('w:hyperlink');link.set(qn('w:anchor'),f'qa_{label}_{i}')
        r=OxmlElement('w:r');text=OxmlElement('w:t');text.text=cap.text;r.append(text);link.append(r);para._p.append(link)
        para.add_run('\t')
        for kind in ['begin','instruction','separate','value','end']:
            r=OxmlElement('w:r')
            if kind=='instruction':el=OxmlElement('w:instrText');el.set(qn('xml:space'),'preserve');el.text=f' PAGEREF qa_{label}_{i} \\h '
            elif kind=='value':el=OxmlElement('w:t');el.text='0'
            else:el=OxmlElement('w:fldChar');el.set(qn('w:fldCharType'),kind)
            r.append(el);para._p.append(r)
d.save(p)
print('caption references',len(d.element.xpath('.//w:bookmarkStart[starts-with(@w:name,"qa_")]')))
