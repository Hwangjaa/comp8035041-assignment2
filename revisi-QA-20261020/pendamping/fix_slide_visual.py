from pathlib import Path
import json
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
root = Path(__file__).parent
src = root / 'Assignment-II-DataCo-Hasil-Aktual-QA.pptx'
out = root / 'Assignment-II-DataCo-Hasil-Aktual-QA-Visual.pptx'
p = Presentation(src)
for slide in p.slides:
    footer = slide.shapes[2]
    footer.top = Inches(6.93)
    footer.height = Inches(.4)
    for para in footer.text_frame.paragraphs:
        for run in para.runs:
            run.font.size = Pt(14)
            run.font.color.rgb = RGBColor.from_string('334155')
slide = p.slides[4]
body = slide.shapes[1]
body.text = 'Semua: precision=0.576402; recall=1.000000\nSemua: TN=0; FP=3.407; FN=0; TP=4.636'
body.top = Inches(4.65)
body.height = Inches(1.5)
for para in body.text_frame.paragraphs:
    for run in para.runs:
        run.font.size = Pt(23)
        run.font.color.rgb = RGBColor.from_string('172554')
rows = [['Model','AP','ROC-AUC','F1'],['baseline','0.576402','0.500000','0.731288'],['sklearn_lr','0.736424','0.719782','0.731288'],['spark_lr','0.736851','0.720601','0.731288']]
table = slide.shapes.add_table(4,4,Inches(.85),Inches(1.85),Inches(11.6),Inches(2.35)).table
for i,row in enumerate(rows):
    for j,value in enumerate(row):
        cell = table.cell(i,j)
        cell.text = value
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor.from_string('172554' if i==0 else 'E8EEF5')
        for para in cell.text_frame.paragraphs:
            for run in para.runs:
                run.font.size = Pt(23)
                run.font.bold = i==0
                run.font.color.rgb = RGBColor.from_string('FFFFFF' if i==0 else '172554')
slide = p.slides[5]
old = slide.shapes[1].text
slide.notes_slide.notes_text_frame.text += '\n\nThreshold presisi penuh (tidak berubah):\n' + old
slide.shapes[1].text = 'H1: maksimum F1 validation; ties → threshold terbesar\nBaseline: 0.572472\nscikit-learn: 0.307476\nSpark: 0.446602\nTampilan 6 desimal; presisi penuh ada pada catatan dan artefak.\nDibekukan sebelum test; prediksi positif p ≥ threshold'
for para in slide.shapes[1].text_frame.paragraphs:
    for run in para.runs:
        run.font.size = Pt(23)
        run.font.color.rgb = RGBColor.from_string('172554')
p.save(out)
q = Presentation(out)
assert len(q.slides)==9
assert [[c.text for c in r.cells] for r in q.slides[4].shapes[-1].table.rows] == rows
assert '0.3074755313277422' in q.slides[5].notes_slide.notes_text_frame.text
result = {'source':str(src),'output':str(out),'slides':9,'metric_table_exact':True,'full_precision_thresholds_preserved_in_notes':True,'changes':['Footer 14 pt dengan kontras lebih jelas','Tabel metrik menggunakan kolom asli','Threshold tampilan 6 desimal; presisi penuh dipertahankan pada notes'],'visual_review':'PENDING'}
(root/'slide-visual-edit.json').write_text(json.dumps(result,indent=2),encoding='utf-8')
print(json.dumps(result,indent=2))
