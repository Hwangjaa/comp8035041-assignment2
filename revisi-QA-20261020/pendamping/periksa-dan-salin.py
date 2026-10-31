from pathlib import Path
import json, hashlib, math
import pandas as pd
import networkx as nx
from sklearn.metrics import average_precision_score, roc_auc_score, precision_score, recall_score, f1_score, confusion_matrix
from pptx import Presentation
P=Path('/Users/rafaelyudianto/Documents/COMP8035041-Assignment-II'); O=Path(__file__).parent
R=P/'eksekusi-final-lokal/run-actual';D=P/'eksekusi-final-lokal/deliverables'
checks=[]
def check(n,a,b):
 ok=math.isclose(float(a),float(b),rel_tol=1e-10,abs_tol=1e-12) if isinstance(a,(int,float)) else a==b
 checks.append({'pemeriksaan':n,'sesuai':bool(ok),'aktual':a,'acuan':b})
 assert ok,(n,a,b)
metrics=json.loads((R/'metrics.json').read_text())
for m in metrics:
 f=pd.read_csv(R/f"predictions-{m['partition']}-{m['model']}.csv")
 y=f['target'] if 'target' in f else f['Late_delivery_risk']
 pred=f['prediction'];prob=f['probability_1']
 check(f"{m['model']} {m['partition']} n",len(f),m['n'])
 for k,fn in [('average_precision',average_precision_score),('roc_auc',roc_auc_score)]:check(f"{m['model']} {m['partition']} {k}",float(fn(y,prob)),m[k])
 for k,fn in [('precision',precision_score),('recall',recall_score),('binary_f1',f1_score)]:check(f"{m['model']} {m['partition']} {k}",float(fn(y,pred,zero_division=0)),m[k])
 check(f"{m['model']} {m['partition']} CM",confusion_matrix(y,pred,labels=[0,1]).ravel().tolist(),m['confusion_matrix'])
 check(f"{m['model']} {m['partition']} threshold",bool(((prob>=m['threshold']).astype(int)==pred).all()),True)
edges=pd.read_csv(R/'graph-edges.csv');print('kolom edge',edges.columns.tolist())
G=nx.from_pandas_edgelist(edges,source='source',target='target')
g=json.loads((R/'graph-summary.json').read_text())
check('graph nodes',G.number_of_nodes(),g['nodes']);check('graph edges',G.number_of_edges(),g['edges']);check('graph components',nx.number_connected_components(G),g['connected_components']);check('self loops',nx.number_of_selfloops(G),0)
check('duplicate edges',int(edges.duplicated().sum()),0)
for cid,d in [(17,19534),(18,17954),(24,17088)]:check(f'degree Category Id {cid}',G.degree(f'category:{cid}'),d)
prs=Presentation(D/'Assignment-II-DataCo-Hasil-Aktual.pptx')
script=(D/'Naskah-Presentasi-Faktual.md').read_text()
old='Video teknis diproduksi terpisah; draf teknis AI bukan final mahasiswa.'
new='Produksi video dibatalkan sesuai instruksi pengguna. Brief tetap meminta video presentasi maksimal 10 menit; pemenuhan atau pengecualian belum terbukti. Draf ini bukan video final.'
for s in prs.slides:
 if s.has_notes_slide:
  for para in s.notes_slide.notes_text_frame.paragraphs:
   for run in para.runs:
    if old in run.text:run.text=run.text.replace(old,new)
script=script.replace(old,new).replace('- Tinjau video teknis yang diproduksi terpisah; draf teknis AI bukan final mahasiswa. Lengkapi segmen mahasiswa dan tautan final.','- Produksi video dibatalkan sesuai instruksi pengguna; tidak ada video dibuat dalam audit ini. Brief tetap mensyaratkan video maksimal 10 menit. Mahasiswa perlu memastikan pemenuhan atau memperoleh pengecualian resmi; statusnya belum terbukti.')
# Koreksi status tampak pada slide tanpa menambah klaim hasil eksperimen.
for sh in prs.slides[8].shapes:
 if sh.has_text_frame:
  for para in sh.text_frame.paragraphs:
   for run in para.runs:
    run.text=run.text.replace('Tautan berbagi, deklarasi AI dan video diselesaikan penulis','Sharing/deklarasi belum final; produksi video dibatalkan')
out=O/'Assignment-II-DataCo-Hasil-Aktual-QA.pptx';prs.save(out)
(O/'Naskah-Presentasi-Faktual-QA.md').write_text(script)
re=Presentation(out)
check('jumlah slide salinan',len(re.slides),9)
check('catatan video dikoreksi',new in re.slides[8].notes_slide.notes_text_frame.text,True)
check('naskah tanpa klaim video diproduksi',old not in script,True)
for i in range(8):check(f'teks slide {i+1} tetap',[s.text for s in re.slides[i].shapes if s.has_text_frame],[s.text for s in Presentation(D/'Assignment-II-DataCo-Hasil-Aktual.pptx').slides[i].shapes if s.has_text_frame])
outline=[{'slide':i+1,'teks':[s.text for s in sl.shapes if s.has_text_frame],'catatan':sl.notes_slide.notes_text_frame.text} for i,sl in enumerate(re.slides)]
(O/'outline-slide-QA.json').write_text(json.dumps(outline,ensure_ascii=False,indent=2))
orig=json.loads((O/'hash-sumber-awal.json').read_text())
for file,h in orig.items():check('sumber tidak berubah: '+Path(file).name,hashlib.sha256(Path(file).read_bytes()).hexdigest(),h)
(O/'verifikasi-pendamping.json').write_text(json.dumps({'seluruh_pemeriksaan_lulus':True,'jumlah_pemeriksaan':len(checks),'pemeriksaan':checks,'render_visual':'Tidak dilakukan: soffice/pdftoppm tidak tersedia. Verifikasi struktur/teks dan pembukaan ulang PPTX berhasil.'},ensure_ascii=False,indent=2))
print(json.dumps({'jumlah_pemeriksaan':len(checks),'lulus':True,'salinan':str(out)},ensure_ascii=False))
