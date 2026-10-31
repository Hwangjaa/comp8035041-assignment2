"""Revisi salinan laporan; tidak mengubah hasil atau metode eksperimen."""
from pathlib import Path
from io import BytesIO
import json,re,hashlib,copy
from docx import Document
from docx.shared import Inches,Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch
ROOT=Path(__file__).resolve().parents[1]
OUT=Path(__file__).resolve().parent
SRC=ROOT/'eksekusi-final-lokal/deliverables/Assignment-II-DataCo-Laporan-Hasil-Aktual-Italic-Corrected.docx'
RUN=ROOT/'eksekusi-final-lokal/run-actual'
DEST=OUT/'Assignment-II-DataCo-Revisi-QA-v2.docx'
OUT.mkdir(exist_ok=True)
BEFORE=hashlib.sha256(SRC.read_bytes()).hexdigest()
d=Document(SRC)
math_before=[__import__('lxml.etree',fromlist=['tostring']).tostring(x) for x in d.element.xpath('.//m:oMath')]
changes=[]
def set_text(p,text):
    assert not p._p.xpath('.//w:drawing|.//m:oMath|.//w:fldChar|.//w:instrText')
    p.clear();p.add_run(text).italic=False
    changes.append({'new':text})
def insert_after(anchor,text='',style=None):
    p=d.add_paragraph(text,style=style)
    anchor._p.addnext(p._p)
    return p

def diagram(file,title,panels):
    fig,ax=plt.subplots(figsize=(12,7.5));ax.set_xlim(0,12);ax.set_ylim(0,7.5);ax.axis('off')
    ax.text(.25,7.1,title,fontsize=15,weight='bold',color='#163047')
    colors=['#e7edf4','#e7f3e9','#fff0d7']
    for i,(head,body) in enumerate(panels):
        y=6.55-i*2.05
        ax.add_patch(FancyBboxPatch((.25,y-1.6),11.5,1.65,boxstyle='round,pad=0.08',facecolor=colors[i],edgecolor='#8193a3'))
        ax.text(.5,y-.25,head,fontsize=12,weight='bold',color='#163047')
        ax.text(.5,y-.65,body,fontsize=10.5,va='top',linespacing=1.6,color='#172b3a')
    fig.savefig(OUT/file,dpi=180,bbox_inches='tight');plt.close(fig)
diagram('gambar-2.1-arsitektur-revisi.png','Arsitektur konseptual, historis, dan implementasi aktual',[
('KONSEPTUAL — rancangan Assignment I','Hadoop / HDFS → Spark → analitik → konsumsi operasional\nLapisan data lake, Parquet/JSON, dan streaming merupakan desain; bukan bukti implementasi.'),
('HISTORIS — bukti Assignment I','Snapshot DataCo → audit / pembersihan lokal; pemeriksaan SHA-256\nHDFS single-node hanya menulis / membaca /smoke/input.txt; dataset lengkap tidak dimuat.'),
('TERLAKSANA — Assignment II lokal','CSV file:/// → pesanan / split D2 / fitur training-only → LR scikit-learn + Spark local[2]\nPrediksi / evaluasi H1 tersimpan; edge Spark distinct → degree / komponen NetworkX satu mesin.')])
diagram('gambar-2.2-alur-revisi.png','Alur eksekusi aktual dan batas pengujian',[
('1. DATA DAN TRANSFORMASI — terlaksana','CSV mentah → audit nonbatal / singleton / target → satu baris per Order Id\nSplit temporal D2 → vocabulary fit hanya training → matriks identik, 48 fitur, tanpa scaling.'),
('2. MODEL DAN GRAPH — terlaksana','Baseline prior + LR scikit-learn + LR Spark → threshold H1 validation → prediksi test / metrik\nSeluruh pasangan unik order–kategori → graph lengkap; subset 210 node / 495 edge hanya tampilan.'),
('3. BATAS — tidak diuji','Tidak ada pemuatan dataset lengkap ke HDFS, streaming, benchmark, atau cluster multi-node.\nKesiapan as-of, manfaat operasional, intervensi, dan biaya bisnis belum dibuktikan.')])
# Ganti dua gambar aktif melalui relasi baru; ukuran tetap sesuai area cetak.
for shape,file in zip(list(d.inline_shapes),['gambar-2.1-arsitektur-revisi.png','gambar-2.2-alur-revisi.png']):
    rid,image=d.part.get_or_add_image(str(OUT/file))
    shape._inline.xpath('.//a:blip')[0].set(qn('r:embed'),rid)
    shape.width=Inches(6.4);shape.height=int(shape.width*image.px_height/image.px_width)
# Perbaikan teks faktual; instruksi editor dipindahkan ke pendamping.
for p in list(d.paragraphs):
    t=p.text
    if t=='[Tanggal pengumpulan ditetapkan penulis]':set_text(p,'Tanggal pengumpulan: belum ditetapkan mahasiswa')
    elif t.startswith('Gambar 2.1 dan Gambar 2.2 dipertahankan'):set_text(p,'Gambar 2.1 membedakan arsitektur konseptual Assignment I, bukti historis, dan implementasi Assignment II yang telah dijalankan. Gambar 2.2 memperlihatkan alur aktual lokal file:/// serta komponen yang tidak diuji. Kedua diagram tidak menyatakan dataset lengkap dimuat ke HDFS.')
    elif t=='Gambar 2.1 Arsitektur konseptual Hadoop–HDFS–Spark dan status implementasi':set_text(p,'Gambar 2.1 Arsitektur konseptual, bukti historis, dan implementasi aktual')
    elif t=='Gambar 2.2 Alur pemrosesan HDFS–Spark dan batas bukti eksekusi':set_text(p,'Gambar 2.2 Alur eksekusi lokal Assignment II dan batas pengujian')
    elif '[WAJIB MAHASISWA:' in t:set_text(p,re.sub(r'\s*\[WAJIB MAHASISWA:.*?\]','',t))
    elif t.startswith('[WAJIB DITULIS MAHASISWA:'):set_text(p,'Bagian analisis mahasiswa belum diselesaikan pada versi tinjauan ini.')
    elif t.startswith('Dataset: https://doi.org/'):
        set_text(p,'Dataset: https://doi.org/10.17632/8gx2fvg2k6.5. Notebook tereksekusi: eksekusi-final-lokal/Assignment-II-DataCo-executed-local.ipynb; kode: eksekusi-final-lokal/src/pipeline.py; hasil: eksekusi-final-lokal/run-actual. Slide dan naskah berada pada eksekusi-final-lokal/deliverables. Tautan berbagi kode dan slide untuk penilai belum diverifikasi; video presentasi belum tersedia. Rincian kelengkapan tercatat pada lembar QA terpisah.')
    elif t.startswith('Assignment Type: ['):
        set_text(p,'Assignment Type: penggunaan AI parsial (Partial); nomor Type pada formulir memerlukan konfirmasi klasifikasi resmi.')
    elif 'Istilah dispatch dalam judul tetap memerlukan keputusan penulis' in t:
        set_text(p,t.replace('Istilah dispatch dalam judul tetap memerlukan keputusan penulis berdasarkan batas interpretasi tersebut.','Judul laporan menggunakan risiko keterlambatan pengiriman; alias dispatch pada jejak historis tidak memperluas makna label sumber.'))
# Daftar memakai style caption khusus, bukan nama label bahasa Inggris yang tidak cocok.
for p in d.element.xpath('.//w:p'):
    for inst in p.xpath('.//w:instrText'):
        if 'TOC' in (inst.text or '') and '\\c' in inst.text:
            is_table='Table' in inst.text or 'Tabel' in inst.text
            inst.text=' TOC \\t "'+('Table Caption' if is_table else 'Figure Caption')+',1" \\h \\z '
# Cakup juga field simple (daftar lama merupakan SDT/fldSimple).
for inst in d.element.xpath('.//w:instrText'):
    if 'TOC' in (inst.text or '') and ('Table Caption' in inst.text or 'Figure Caption' in inst.text):pass
for fld in d.element.xpath('.//w:fldSimple'):
    ins=fld.get(qn('w:instr'),'')
    if 'TOC' in ins and '\\c' in ins:
        fld.set(qn('w:instr'),' TOC \\t "'+('Table Caption' if 'Table' in ins or 'Tabel' in ins else 'Figure Caption')+',1" \\h \\z ')
# Ubah fldSimple daftar menjadi field kompleks agar Word menghitung daftar style khusus.
for fld in list(d.element.xpath('.//w:fldSimple')):
    ins=fld.get(qn('w:instr'),'')
    if 'TOC' not in ins:continue
    parent=fld.getparent();idx=parent.index(fld)
    nodes=[]
    for kind in ['begin','instruction','separate','end']:
        r=OxmlElement('w:r')
        if kind=='instruction':
            v=OxmlElement('w:instrText');v.set(qn('xml:space'),'preserve');v.text=ins
        else:
            v=OxmlElement('w:fldChar');v.set(qn('w:fldCharType'),kind)
        r.append(v);nodes.append(r)
    parent.remove(fld)
    for node in nodes:parent.insert(idx,node);idx+=1
# Tambahkan gambar graph aktual ke Bab 6.1 sebelum Tabel 6.2.
anchor=next(p for p in d.paragraphs if p.text.startswith('Degree kategori menghitung'))
ref=insert_after(anchor,'Gambar 6.1 menampilkan subset graph berdasarkan aturan pemilihan yang sama dengan graph-summary.json. Angka pada Tabel 6.2 dan Tabel 6.3 tetap berasal dari graph penuh, bukan subset gambar.')
figp=insert_after(ref);figp.add_run().add_picture(str(RUN/'graph-subset.png'),width=Inches(6.4));figp.paragraph_format.keep_with_next=True
cap=insert_after(figp,'Gambar 6.1 Subset graph pesanan–kategori: 210 node dan 495 edge','Figure Caption')
cap.paragraph_format.keep_with_next=False
# Tabel validation 7 kolom diubah orientasinya untuk keterbacaan.
m=json.loads((RUN/'metrics.json').read_text());val={r['model']:r for r in m if r['partition']=='validation'}
tab=d.tables[16]
new=d.add_table(rows=8,cols=4);new.style=tab.style
rows=[['Ukuran','Baseline','LR scikit-learn','LR Spark MLlib']]
for label,key in [('Average precision','average_precision'),('ROC-AUC','roc_auc'),('Precision kelas 1','precision'),('Recall kelas 1','recall'),('Binary F1','binary_f1')]:
    rows.append([label]+[f'{val[x][key]:.6f}' for x in ['baseline','sklearn_lr','spark_lr']])
rows.append(['Confusion matrix (TN, FP, FN, TP)']+[', '.join(map(str,val[x]['confusion_matrix'])) for x in ['baseline','sklearn_lr','spark_lr']])
rows.append(['Threshold validation']+[f'{val[x]["threshold"]:.6f}' for x in ['baseline','sklearn_lr','spark_lr']])
for r,values in zip(new.rows,rows):
    for c,t in zip(r.cells,values):c.text=t
old=tab._tbl;old.addprevious(new._tbl);old.getparent().remove(old)
# Penomoran caption yang ada tetap dipertahankan; style menjadi sumber daftar dinamis.
for p in d.paragraphs:
    if p.style.name=='Table Caption':p.paragraph_format.keep_with_next=True
    if p.style.name.startswith('Heading'):
        p.paragraph_format.keep_with_next=True
        if p.style.name=='Heading 1':p.paragraph_format.page_break_before=True
# Hindari row splitting dan ulang header; font tabel minimum 9 pt.
for table in d.tables:
    table.autofit=False
    for ri,row in enumerate(table.rows):
        props=row._tr.get_or_add_trPr()
        if not props.xpath('./w:cantSplit'):props.append(OxmlElement('w:cantSplit'))
        if ri==0 and not props.xpath('./w:tblHeader'):props.append(OxmlElement('w:tblHeader'))
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.keep_with_next=False
                p.paragraph_format.space_after=Pt(3)
                p.paragraph_format.space_before=Pt(3)
                for r in p.runs:
                    r.font.size=Pt(9)
                    if ri==0:r.bold=True
# Beri tahu Word agar field diperbarui.
settings=d.settings.element
u=settings.find(qn('w:updateFields'))
if u is None:u=OxmlElement('w:updateFields');settings.append(u)
u.set(qn('w:val'),'true')
assert [__import__('lxml.etree',fromlist=['tostring']).tostring(x) for x in d.element.xpath('.//m:oMath')]==math_before
assert len(d.inline_shapes)==3
assert hashlib.sha256(SRC.read_bytes()).hexdigest()==BEFORE
d.save(DEST)
# Baca struktur salinan, bukan mengklaim render dari XML.
reopened=Document(DEST)
assert len(reopened.inline_shapes)==3
assert len(reopened.element.xpath('.//m:oMath'))==len(math_before)
(OUT/'report-edit-audit.json').write_text(json.dumps({'source':str(SRC),'source_sha256':BEFORE,'output':str(DEST),'changes':changes,'active_pictures':len(reopened.inline_shapes),'native_math_count':len(math_before),'tables':len(reopened.tables),'method_changed':False},indent=2,ensure_ascii=False))
print(json.dumps({'output':str(DEST),'pictures':len(reopened.inline_shapes),'math':len(math_before),'changes':len(changes)},indent=2))
