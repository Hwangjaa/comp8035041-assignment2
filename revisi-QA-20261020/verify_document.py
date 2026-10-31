from pathlib import Path
import re,json,hashlib
from docx import Document
import pymupdf as fitz
OUT=Path(__file__).resolve().parent
p=OUT/'Assignment-II-DataCo-Revisi-QA-Tinjauan.docx'
d=Document(p);pdf=fitz.open(OUT/'Assignment-II-DataCo-Revisi-QA-Tinjauan.pdf')
texts=[page.get_text() for page in pdf];alltext='\n'.join(texts)
assert len(pdf)<=40 and len(d.inline_shapes)==3 and len(d.element.xpath('.//m:oMath'))==5
assert 'No table of contents entries found' not in alltext
assert 'WAJIB MAHASISWA' not in alltext and 'WAJIB DITULIS' not in alltext
captions=[x.text for x in d.paragraphs if x.style.name in ['Table Caption','Figure Caption']]
# DOCX hasil LO mengganti nama style; judul aktual diambil langsung paragraf tubuh.
if not captions:captions=[x.text for x in d.paragraphs if re.match(r'^(Tabel|Gambar) \d+\.\d+ ',x.text) and '\t' not in x.text]
checks=[]
for caption in captions:
    compact=lambda s:re.sub(r'\s+','',s)
    matches=[i+1 for i,t in enumerate(texts) if compact(caption) in compact(t)]
    assert len(matches)>=2,(caption,matches)
    prefix=re.match(r'^(Tabel|Gambar) \d+\.\d+',caption).group()
    listpage=3 if prefix.startswith('Tabel') else 4
    lines=texts[listpage].splitlines()
    line=next((l for l in lines if prefix in l),None)
    assert line,(caption,'missing list line')
    target=int(re.search(r'(\d+)\s*$',line).group(1))
    assert target in matches and target!=listpage+1,(caption,target,matches)
    checks.append({'caption':caption,'page':target,'list_page':listpage+1,'target_verified':True})
assert len(checks)==23,len(checks)
chapters=[]
for i,text in enumerate(texts):
    lines=text.splitlines()
    for line in lines:
        if re.match(r'^BAB [1-6]\. ',line) and not re.search(r'\.{3}|\d\s*$',line):
            chapters.append({'heading':line,'page':i+1,'at_page_start':lines.index(line)<=2})
assert len(chapters)==6 and all(c['at_page_start'] for c in chapters),chapters
result={'docx':str(p),'pdf_pages':len(pdf),'active_images':len(d.inline_shapes),'native_equations':5,'caption_count':len(checks),'caption_pages':checks,'chapters':chapters,'package_validation':'PASS separate command','visual_review':'BLOCKED: vision_analyze returned HTTP 502 and 504; no claim of visual QA PASS','sha256':hashlib.sha256(p.read_bytes()).hexdigest()}
(OUT/'document-verification.json').write_text(json.dumps(result,indent=2,ensure_ascii=False));print(json.dumps(result,indent=2,ensure_ascii=False))
